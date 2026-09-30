import re, os, math, unicodedata, zipfile
from openpyxl import load_workbook
from openpyxl.styles import Font
from pathlib import Path
HERE=Path(__file__).resolve().parent; WORK=HERE/'work'; INPUT=HERE/'input'
exec(open(HERE/'overrides.py').read())
SRC=str(INPUT/'Inventario_MonteCastelo_JoaoPaulo_Filipinho_v4_GERAL_1.xlsx')
TPL=str(WORK/'template.xlsx')
OUTD=str(WORK/'fichas'); os.makedirs(OUTD,exist_ok=True)
inv=load_workbook(SRC)
H=['n','nome','autor','data','logr','num','bairro','cep','funcao','obs','ref','lat','lon','aba','origem']
recs=[dict(zip(H,r)) for r in inv['Consolidado'].iter_rows(min_row=2,values_only=True) if r[0]]
bib={r[0].strip():r[1] for r in inv['Bibliografia'].iter_rows(min_row=3,values_only=True) if r[0] and r[1]}
# nº por aba
abanum={}
for sh in ['Monte Castelo','João Paulo','Filipinho']:
    for r in inv[sh].iter_rows(min_row=2,values_only=True):
        if r[0]: abanum[(sh,r[1],str(r[5]))]=r[0]
PFX={'Monte Castelo':'MC','João Paulo':'JP','Filipinho':'FL'}
FUN={'Educacional':'Educacional','Residencial':'Residencial','Institucional':'Institucional','Comercial':'Comercial','Industrial':'Industrial','Religioso':'Religiosa','Esportivo':'Esportiva','Institucional (saúde)':'Saúde','Comercial/Institucional':'Comercial','Residencial/Comercial':'Residencial'}
LOCAIS=["Retiro Natal","Fátima","Vila Passos","Belira","Canto da Fabril","Apeadouro","Alemanha","Coreia","Vila Ivar Saldanha","João Paulo","Filipinho"]
CONSMAP={'boa':'Bom','regular':'Regular','ruim':'Ruim'}
INV="Inventário v4"
def expand(l):
    if not l: return l
    return re.sub(r'^Av\.\s','Avenida ',re.sub(r'^R\.\s','Rua ',l))
def slug(s,n=45):
    s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode()
    s=re.sub(r'[^A-Za-z0-9]+','_',s).strip('_'); return s[:n].rstrip('_')
def year_of(d):
    if not d: return None
    m=re.search(r'(1[89]\d\d|20\d\d)',str(d)); return int(m.group(1)) if m else None
def periodo(y):
    if y is None: return None
    if 1930<=y<=1939: return "1930–1939"
    if 1940<=y<=1949: return "1940–1949"
    if 1950<=y<=1959: return "1950–1959"
    if 1960<=y<=1969: return "1960–1969"
    if 1970<=y<=1979: return "1970–1979"
    if y==1980: return "1980"
    return None
GENERIC=re.compile(r'^(Residência( da | de | \(|$)(?!família)|Casa da |Escola$|Posto de Gasolina|Vila Militar)')
NEWSP=['Imparcial','Correio','Pacotilha','Jornal','Novidades','O Dia']

def build_index(ws):
    idx={}; opts={}; last=ws.max_row
    for r in range(6,last+1):
        a=ws.cell(r,1).value
        if a is not None and str(a) not in idx: idx[str(a)]=r
    return idx
def rows_after(ws,start,maxn=60):
    out=[]; r=start+1
    while r<ws.max_row and len(out)<maxn:
        b=ws.cell(r,2).value
        if isinstance(b,str) and b.startswith("   ☐  "): out.append((b[6:],r)); r+=1
        else: break
    return out
def find_text_row(ws,text):
    for r in range(6,ws.max_row+1):
        v=ws.cell(r,1).value
        if isinstance(v,str) and v.startswith(text): return r
    raise KeyError(text)
def table_rows(ws,anchor,labels_needed=None,maxscan=25):
    s=find_text_row(ws,anchor); out={}
    for r in range(s+1,s+maxscan):
        b=ws.cell(r,2).value
        if b is not None: out.setdefault(str(b),r)
    return out,s

def fit(ws,r):
    mx=0
    for c,w in ((4,46),(5,30)):
        v=ws.cell(r,c).value
        if v is None: continue
        t=str(v); lines=sum(max(1,math.ceil(len(p)/(w*1.15))) for p in t.split("\n"))
        mx=max(mx,lines)
    need=min(409,12*mx+6)
    cur=ws.row_dimensions[r].height or 18
    if need>cur: ws.row_dimensions[r].height=need

stats=[]
for rec in recs:
    n=rec['n']; o=O.get(n,{})
    wb=load_workbook(TPL); ws=wb['FIAMS']; idx=build_index(ws)
    filled=[0]
    def put(code,val,src=None,col=4):
        if val in (None,""): return
        r=idx[code]; ws.cell(r,col,val)
        if src: ws.cell(r,5,src)
        filled[0]+=1; fit(ws,r)
    def srcnote(code,txt):
        r=idx[code]; cur=ws.cell(r,5).value
        ws.cell(r,5,(cur+" · " if cur else "")+txt); fit(ws,r)
    def ov(code,default=None,dsrc=None):
        v=o.get(code,default)
        if v is None: return
        if isinstance(v,tuple): put(code,v[0],v[1])
        else: put(code,v,dsrc if code not in o else (dsrc or "OBS. — "+INV))
    obs=rec['obs'] or ""; ref=rec['ref'] or ""; nome=rec['nome'] or ""
    aba=rec['aba']; pfx=PFX[aba]; code=f"FIAMS-{pfx}-{n:03d}"
    ab=abanum.get((aba,rec['nome'],str(rec['num'])))
    alerts=[]
    # 1.1
    put("1.1.1",code,"Atribuído: prefixo do bairro-aba + nº do Consolidado (Inventário v4)")
    put("1.1.2",f"Consolidado nº {n}"+(f" · aba {aba} nº {ab}" if ab else f" · aba {aba}")+f" · origem: {rec['origem']}",INV)
    put("1.1.6","v0.1 — pré-preenchida a partir do Inventário v4","Revisar após visita de campo")
    # 1.2 denominações
    generic=bool(GENERIC.match(nome))
    if not any(k in o for k in ("1.2.1","1.2.2")):
        if generic: put("1.2.4",nome,"Denominação descritiva do "+INV)
        elif rec['origem']=="novo (2026)": put("1.2.1",nome,"Denominação no levantamento geocodificado de 2026 ("+INV+")")
        else: put("1.2.2",nome,"EDIFICAÇÃO — "+INV)
    else:
        for c in ("1.2.1","1.2.2"): ov(c)
        if not generic and nome not in [str(ws.cell(idx[c],4).value) for c in ("1.2.1","1.2.2")]:
            o.setdefault("1.2.4",nome+" (denominação no "+INV+")")
    for c in ("1.2.3","1.2.4","1.2.5"): ov(c)
    # 1.3
    logr=expand(rec['logr']); put("1.3.1",logr,"LOGRADOURO — "+INV)
    put("1.3.2",str(rec['num']) if rec['num'] else None,"Nº END. — "+INV)
    ov("1.3.3")
    b=rec['bairro']; put("1.3.4",aba,"Aba do "+INV+" (recorte: Monte Castelo e subáreas, João Paulo, Filipinho)")
    loc={'Canto Fabril':'Canto da Fabril'}.get(b,b)
    if aba=='Monte Castelo' and loc!='Monte Castelo':
        if loc in LOCAIS: put("1.3.5",loc,"BAIRRO no "+INV)
        else: put("1.3.5","Outra","BAIRRO no "+INV+f": {b} — fora da lista de subáreas do recorte")
    if b not in (aba,) and loc not in LOCAIS: pass
    put("1.3.7",rec['cep'],"CEP — "+INV)
    if rec['cep'] in ("65137-000","65000-000"): srcnote("1.3.7","CEP A CONFERIR (fora da faixa do recorte / genérico)"); alerts.append("CEP a conferir")
    if not rec['cep']: alerts.append("sem CEP")
    fonte_end="Revisão de endereços e geocodificação (2026) — "+INV if rec['lat'] else "Inventário v3 (sem geocodificação em 2026)"
    if "base cartográfica" in obs: fonte_end+="; base cartográfica"
    put("1.3.9",fonte_end)
    AF=("divergência de endereço","divergência de bairro","divergência de logradouro","registra o endereço","não localizado","identificação incerta","sem geocodificação","numeração a confirmar","logradouro a conferir","endereço e bairro a conferir","verificar bairro","confirmação em campo recomendada","confirmação em campo indispensável","divergência de numeração")
    flags=[k for k in AF if k in obs.lower()]
    conf="Baixa" if flags or not rec['lat'] else "Média"
    put("1.3.10",conf,"Preliminar (dados documentais, sem confirmação em campo)"+(f" — pendências no OBS.: {', '.join(sorted(set(f.lower() for f in flags)))}" if flags else ""))
    for c in ("1.3.1E","1.3.2E","1.3.5E"):
        if c in o: srcnote(c[:-1],o[c])
    if re.search(r'CORREÇÃO DE (ENDEREÇO|BAIRRO)',obs) and "1.3.1E" not in o:
        m=re.search(r'CORREÇÃO DE (?:ENDEREÇO|BAIRRO) \(2026\):([^.]*(?:\.[^A-Z][^.]*)*)',obs)
        srcnote("1.3.1","CORREÇÃO (2026):"+(m.group(1) if m else " ver OBS."))
    if any(k in obs for k in ("DIVERGÊNCIA DE ENDEREÇO","DIVERGÊNCIA DE BAIRRO","DIVERGÊNCIA DE LOGRADOURO","DIVERGÊNCIA: o artigo")): alerts.append("divergência de endereço")
    # 1.4
    if rec['lat']:
        put("1.4.1",float(rec['lat']),"LATITUDE — geocodificação (2026), "+INV+" — substituir pela leitura de GPS em campo")
        put("1.4.2",float(rec['lon']),"LONGITUDE — geocodificação (2026), "+INV)
        put("1.4.8","Base cartográfica / imagem de satélite","Geocodificação (2026)")
    else:
        alerts.append("sem coordenadas"); srcnote("1.4.1","SEM GEOCODIFICAÇÃO no "+INV+" — coletar em campo")
    # 1.5 reconhecimento
    rows=dict(rows_after(ws,idx["1.5.2"]))
    acad=[k for k in ("TRINDADE E SILVA; LOPES, 2026","TRINDADE E SILVA, 2024","CHIQUITELLI, 2020") if k in ref]
    if acad:
        ws.cell(rows["Inventário acadêmico"],4,"Sim"); ws.cell(rows["Inventário acadêmico"],5,"; ".join(acad)); filled[0]+=1
    pubs=[k for k in bib if k in ref and not k.startswith(("@","Fontes","Acervos")) and k not in acad]
    if pubs:
        r_=rows["Publicação especializada"]; ws.cell(r_,4,"Sim"); ws.cell(r_,5,"; ".join(pubs)); fit(ws,r_); filled[0]+=1
    # 2 caracterização
    ov("2.1.1")
    f=FUN.get(rec['funcao'])
    if rec['origem']=="novo (2026)":
        if "2.1.3" not in o and f: put("2.1.3",f,"Inferida da denominação atual (geocodificação 2026) — confirmar em campo")
        alerts.append("função inferida")
    else:
        if "2.1.2" not in o and f: put("2.1.2",f,f"FUNÇÃO no {INV}: {rec['funcao']}")
    ov("2.1.2"); ov("2.1.3")
    if "2.2.1" in o: ov("2.2.1")
    elif rec['funcao']=="Residencial" and re.match(r'^(Residência|Casa)',nome) and rec['origem']!="novo (2026)":
        put("2.2.1","Residência unifamiliar","Inferida da denominação — confirmar")
    ov("2.3.1"); ov("2.3.2",dsrc="OBS. — "+INV)
    # 3 autoria
    a=rec['autor']
    if "3.1.1" in o: ov("3.1.1")
    elif a:
        put("3.1.1","Não identificado" if a.startswith("Autoria não identificada") else a,"ARQUITETO/AUTOR — "+INV)
    if "3.1.1E" in o: srcnote("3.1.1",o["3.1.1E"])
    if "3.1.4" in o: ov("3.1.4")
    elif a:
        if a.startswith("Autoria não identificada") or str(ws.cell(idx["3.1.1"],4).value).startswith("Não identificado"):
            put("3.1.4","Desconhecida",INV)
        else: put("3.1.4","Provável","Autoria registrada em fontes secundárias ("+INV+") — validar")
    if ws.cell(idx["3.1.1"],4).value and ws.cell(idx["3.1.1"],4).value!="Não identificado":
        put("3.1.5",ref,"REFERÊNCIAS — "+INV)
    for c in ("3.1.3","3.2.4","3.2.6","3.3.2","3.3.3","3.4.1","3.4.2","3.4.3"): ov(c)
    # proprietário
    if "3.3.1" in o: ov("3.3.1")
    else:
        m=re.match(r'^Residência (?:Srª|Sr\.|Dr\.) (.+)$',nome)
        if m: put("3.3.1",m.group(1),"Inferido da denominação ("+INV+")")
    if "3.3.4" in o: ov("3.3.4")
    elif ws.cell(idx["3.3.1"],4).value: put("3.3.4","Privada","Residência particular ("+INV+")")
    # obras do mesmo autor
    if "3.4.2" not in o and a and any(x in a for x in ("Cleon Furtado","João Magalhães de Araújo")):
        who="Cleon Furtado" if "Cleon" in a else "João Magalhães de Araújo"
        outros=[f"{x['nome']} (nº {x['n']})" for x in recs if x['n']!=n and x['autor'] and who in x['autor']]
        put("3.4.2",f"No {INV} ({len(outros)}): "+"; ".join(outros),"Registros do mesmo autor no "+INV)
    if "3.4.3" not in o and a and "Cleon Furtado" in a: put("3.4.3","Produção de Cleon Furtado","Pré-classificação pela autoria registrada — validar")
    # 4 cronologia
    trows,_=table_rows(ws,"4. CRONOLOGIA DO BEM")
    cr=dict(o.get("crono",{}))
    m=re.search(r"Imóvel (\d+) do levantamento de 2026[^.]*?:\s*([^.]*)",obs)
    surv=None
    if m:
        item=m.group(1); txt=m.group(2)
        cm=re.search(r'conservação (boa|regular|ruim)',txt); im=re.search(r"IIM '([^']+)'",obs)
        surv={"item":item,"cons":CONSMAP.get(cm.group(1)) if cm else ("Ruína" if "ruína" in txt else None),"iim":im.group(1) if im else None,"txt":txt.strip()}
        if "Situação atual" not in cr: cr["Situação atual"]=(f"2026 — {txt.strip()}"+(f"; IIM '{surv['iim']}'" if surv['iim'] and surv['iim'] not in txt else ""),f"{TS}, Quadro 5, item {item}")
    for ev,(d,s) in cr.items():
        r_=trows[ev]; ws.cell(r_,4,d); ws.cell(r_,5,s); ws.cell(r_,6,"Provável (fonte secundária)"); fit(ws,r_); filled[0]+=1
    dt=rec['data']
    if "4.1" in o: ov("4.1")
    elif dt and str(dt).strip() not in ("-",""): put("4.1",str(dt),"DATA — "+INV)
    y=year_of(o["4.1"][0] if isinstance(o.get("4.1"),tuple) else (o.get("4.1") or dt))
    if y:
        p=periodo(y)
        if p: put("4.2",p,f"Derivado da data registrada ({y})")
        elif y>1980:
            srcnote("4.2",f"Data registrada ({y}) POSTERIOR ao recorte temporal 1930–1980 — verificar critério de inclusão"); alerts.append(f"fora do recorte temporal ({y})")
        elif y<1930: srcnote("4.2",f"Data ({y}) anterior ao recorte")
    # 5
    if obs: put("5.1.1",obs,"Transcrição do campo OBS. do "+INV+" — redigir como síntese crítica; divergências a manter explícitas")
    if "5.2.3" not in o and logr and any(k in logr for k in ("Getúlio Vargas","João Pessoa")):
        put("5.2.3","Imóvel com frente para o eixo Av. Getúlio Vargas – Av. João Pessoa (a João Pessoa é a continuação da Getúlio Vargas em direção ao João Paulo)","Notas e critérios — "+INV)
    ov("5.2.4")
    # 6 fontes
    frows=dict(rows_after(ws,idx["6.1"]))
    f61=dict(o.get("6.1",{}))
    if "Jornal" not in f61:
        seg=[s_.strip() for s_ in ref.split(";") if any(k in s_ for k in NEWSP)]
        if seg: f61["Jornal"]=("Localizada","; ".join(seg))
    for k,(v,s) in f61.items():
        r_=frows[k]; ws.cell(r_,4,v); ws.cell(r_,5,s); fit(ws,r_); filled[0]+=1
    for c in ("6.2.1","6.2.3","6.2.4","6.3.4"): ov(c)
    # 7-9
    for c in ("7.1.1","7.2.1")+tuple(f"7.3.{i}" for i in range(1,12))+("8.1.1","8.1.3","8.1.4","8.1.5","8.1.7","8.1.8","8.2.2","8.2.4","9.1.1","9.1.2","9.2.1","9.3.1","9.4.1","9.4.2","9.5.1","9.5.2"):
        ov(c)
    for chk in ("8.3","9.6","10.1"):
        if chk in o:
            rr=dict(rows_after(ws,idx[chk]))
            for opt in o[chk]:
                ws.cell(rr[opt],4,"Sim"); ws.cell(rr[opt],5,"OBS. — "+INV+" (a verificar em campo)"); filled[0]+=1
    if "9.6E" in o:
        rr=dict(rows_after(ws,idx["9.6"])); ws.cell(rr["Brises"],5,o["9.6E"])
    for c in ("10.2","10.3"): ov(c)
    # 11 uso
    if "11.1" in o: ov("11.1")
    elif rec['origem']!="novo (2026)" and rec['funcao'] and rec['funcao']!="A confirmar": put("11.1",rec['funcao'],"FUNÇÃO — "+INV)
    ov("11.2")
    if "11.2E" in o: srcnote("11.2",o["11.2E"])
    if "usos" in o:
        s_=find_text_row(ws,"11.3"); r0=s_+2
        for i,(per,uso,resp,fo) in enumerate(o["usos"][:4]):
            ws.cell(r0+i,2,per); ws.cell(r0+i,4,uso); ws.cell(r0+i,5,resp); ws.cell(r0+i,6,fo); fit(ws,r0+i); filled[0]+=1
    for c in ("12.2.1","12.2.2","12.2.3"): ov(c)
    # 14.2 IIM do levantamento
    if surv and surv['iim']:
        ws.cell(idx["14.2.3"],5,f"Levantamento 2026: IIM '{surv['iim']}' ({TS}, item {surv['item']}). Os estados A1–A7 não constam do Inventário — preencher em campo."); fit(ws,idx["14.2.3"])
    # 15.2
    if o.get("alt"):
        s_=find_text_row(ws,"15.2"); r0=s_+3
        for i,(t,el,g) in enumerate(o["alt"][:6]):
            ws.cell(r0+i,2,t); ws.cell(r0+i,4,el);
            if g: ws.cell(r0+i,5,g)
            fit(ws,r0+i); filled[0]+=1
    # 16
    if "16.1" in o: ov("16.1")
    elif surv and surv['cons']: put("16.1",surv['cons'],f"Levantamento 2026 ({TS}, item {surv['item']}) — atualizar em campo")
    if "16.1E" in o: srcnote("16.1",o["16.1E"])
    if "pato" in o:
        prow,_=table_rows(ws,"17. MAPA",maxscan=20)
        for ag,(pres,loc) in o["pato"].items():
            ws.cell(prow[ag],4,pres); ws.cell(prow[ag],5,loc); filled[0]+=1
    # 19.10 subsídio
    if "19.10E" in o: ws.cell(idx["19.10"],5,o["19.10E"]); fit(ws,idx["19.10"])
    ov("22.3.2")
    # 24 fontes
    keys=[k for k in bib if k in ref and not k.startswith("Fontes") and not k.startswith("Acervos")]
    acadk=[k for k in keys if re.search(r'Monografia|Dissertação|Tese|Trabalho de Conclusão|TCC',bib[k])]
    digk=[k for k in keys if k.startswith("@")]
    bibk=[k for k in keys if k not in acadk and k not in digk]
    if bibk: put("24.2","\n".join(bib[k] for k in bibk),"Referências completas: aba Bibliografia do "+INV)
    if acadk: put("24.7","\n".join(bib[k] for k in acadk),"Aba Bibliografia do "+INV)
    urls=re.findall(r'https?://\S+',ref)
    dig=[bib[k] for k in digk]+[u.rstrip(';,') for u in urls]
    if dig: put("24.6","\n".join(dig),"Aba Bibliografia / REFERÊNCIAS do "+INV)
    hem=[s_.strip() for s_ in ref.split(";") if any(k in s_ for k in NEWSP)]
    if hem: put("24.1","Fontes hemerográficas: "+"; ".join(hem),"REFERÊNCIAS — "+INV)
    ico=[s_.strip() for s_ in ref.split(";") if re.search(r'[Ff]oto|[Aa]cervo|Álbum',s_)]
    if ico: put("24.4","; ".join(ico),"REFERÊNCIAS — "+INV)
    if "Levantamento de campo e geocodificação" in ref: put("24.3","Levantamento de campo e geocodificação do autor (2026)","REFERÊNCIAS — "+INV)
    if not ref: srcnote("24.2","Registro sem referência no "+INV)
    ws.cell(idx["24.2"],5).value = (ws.cell(idx["24.2"],5).value or "")+("\nREFERÊNCIAS (transcrição): "+ref if ref else ""); fit(ws,idx["24.2"])
    # alertas extra
    if re.search(r'DEMOLID[AO]S? (em|e)\b',obs) or "SUBSTITUÍDA" in obs: alerts.append("demolido/substituído")
    if flags: alerts.append("endereço a confirmar")
    if "não localizado em campo" in obs: alerts.append("não localizado em campo (2026)")
    if "DIVERGÊNCIA DE DATA" in obs or "data a conferir" in obs or re.search("DIVERGÊNCIA|A CONFERIR",str(o.get("4.1",""))): alerts.append("divergência de data")
    # nota de pré-preenchimento no topo (linha 3, col F já usada) -> use A4
    fname=f"{code}_{slug(nome)}.xlsx"
    wb.save(os.path.join(OUTD,fname))
    stats.append(dict(code=code,file=fname,n=n,nome=nome,end=f"{logr}, {rec['num']}" if rec['num'] else logr,bairro=rec['bairro'],aba=aba,autor=a,data=dt,funcao=rec['funcao'],
        surv=(f"item {surv['item']} · {surv['cons'] or '—'} · IIM {surv['iim'] or '—'}" if surv else ""),alerts="; ".join(dict.fromkeys(alerts)),filled=filled[0],origem=rec['origem']))
import json; json.dump(stats,open(WORK/'stats.json','w'),ensure_ascii=False,default=str)
print(len(stats))
