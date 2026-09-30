from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ARIAL="Arial"
def F(sz=9,b=False,c="1A1A1A",i=False,u=None): return Font(name=ARIAL,size=sz,bold=b,color=c,italic=i,underline=u)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color="A6B4BE"); thick_or=Side(style="thick",color="E07B00")
BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
WRAP=Alignment(wrap_text=True,vertical="center")
WRAPT=Alignment(wrap_text=True,vertical="top")
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)

C_TITLE="1F3A4D"; C_SEC="2E5A6E"; C_SUB="DCE6EC"; C_TH="B8CAD6"; C_NOTE="F3F3F3"
CAMPO="CAMPO"; DOC="DOCUMENTAL"; ANA="ANÁLISE"; CD="CAMPO + DOC"
TAGFILL={CAMPO:"E07B00",CD:"E07B00",DOC:"3D6E9E",ANA:"4F7F45"}
AMBER="FFE9C7"; AMBER_L="FFF6EA"

wb=Workbook()
ws0=wb.active; ws0.title="Apresentação"
ws=wb.create_sheet("FIAMS")
wr=wb.create_sheet("Roteiro_Campo")
wm=wb.create_sheet("Monitoramento")
wsig=wb.create_sheet("SIG_Tabela")
wdic=wb.create_sheet("SIG_Dicionário")
wl=wb.create_sheet("Listas")

# ---------------- LISTAS ----------------
LISTS={
 "SNV":["Sim","Não","Não verificável"],
 "SN":["Sim","Não"],
 "DOCG":["Sim","Não","Em elaboração","Não se aplica"],
 "FONTE":["Localizada","Não localizada","Não pesquisada"],
 "BAIRRO":["Monte Castelo","João Paulo","Filipinho","Entorno (fora do recorte)"],
 "LOCAL":["Retiro Natal","Fátima","Vila Passos","Belira","Canto da Fabril","Apeadouro","Alemanha","Coreia","Vila Ivar Saldanha","João Paulo","Filipinho","Outra"],
 "CONF":["Alta","Média","Baixa"],
 "PREC":["GNSS geodésico (< 1 m)","GPS de navegação / celular (± 5–10 m)","Base cartográfica / imagem de satélite","Estimada"],
 "PROT":["Tombado","Entorno de bem tombado","Inserido em sítio tombado","Inventariado","Sem proteção identificada"],
 "NAT":["Edificação isolada","Conjunto de edificações","Edificação integrada a conjunto urbano","Equipamento urbano"],
 "CAT":["Institucional","Residencial","Comercial","Industrial","Educacional","Religiosa","Saúde","Administrativa","Cultural","Esportiva","Outra"],
 "LING":["Arquitetura racionalista","Proto-modernismo","Art Déco / transição","Arquitetura moderna","International Style","Brutalismo","Regionalismo moderno","Modernismo tropical","Modernismo tardio","Hibridismo / transição","Caracterização inconclusiva"],
 "AUT":["Sim","Provável","Atribuição","Desconhecida"],
 "NATC":["Pública","Privada","Religiosa","Empresarial","Outra"],
 "VETOR":["Egressos da Escola Técnica","Obras públicas (DER-MA)","Produção de Cleon Furtado","Projeto importado (escritório externo)","Não identificado","Outro"],
 "PER":["1930–1939","1940–1949","1950–1959","1960–1969","1970–1979","1980"],
 "CERT":["Confirmado (fonte primária)","Provável (fonte secundária)","Estimado (análise formal)","Depoimento oral"],
 "COMP":["Sim","Parcialmente","Não"],
 "FORMA":["Regular","Irregular","Triangular","Trapezoidal","Outra"],
 "IMPL":["Isolada","Encostada em uma divisa","Encostada em duas divisas","Entre medianeiras","Bloco livre","Pavilhonar","Conjunto"],
 "ACESSO":["Acesso total","Acesso parcial","Sem acesso ao interior"],
 "ESTR":["Concreto armado","Estrutura metálica","Alvenaria estrutural","Estrutura mista","Madeira","Outro","Não verificável"],
 "ESQ":["Original","Parcialmente substituída","Substituída","Não verificável"],
 "CONS":["Excelente","Bom","Regular","Ruim","Péssimo","Ruína"],
 "COMPAT":["Compatível com a concepção original","Parcialmente compatível","Incompatível","Não aplicável"],
 "AMB":["Sim","Não"],
 "AUTENT":["Alta","Média","Baixa","Inconclusiva"],
 "INTEG":["Alta","Média","Baixa","Não avaliável"],
 "GRAU":["Alto","Médio","Baixo","Muito baixo"],
 "IIM":["Preservado","Adaptado","Suprimido / Ausente","n.a.","s.d."],
 "INTERV":["Conservação","Manutenção","Reforma","Ampliação","Adaptação","Acréscimo","Supressão","Substituição","Demolição parcial"],
 "REV":["Reversível","Parcialmente reversível","Irreversível"],
 "COMPO":["Bom","Regular","Ruim","Crítico","Não verificável"],
 "PATO":["Presente","Ausente","Não verificável"],
 "INT3":["Baixa","Média","Alta"],
 "TEND":["↑ Crescente","→ Estável","↓ Decrescente"],
 "RISCO":["Baixo","Moderado","Alto","Crítico"],
 "VAL":["Excepcional","Alto","Médio","Baixo"],
 "VAL5":["Excepcional","Muito alto","Alto","Médio","Baixo"],
 "REPR":["Exemplar excepcional","Exemplar muito representativo","Exemplar representativo","Exemplar secundário","Representatividade inconclusiva"],
 "RAR":["Muito rara","Rara","Comum","Muito comum"],
 "AMA":["Alta","Média","Baixa"],
 "HU":["Muito alta","Alta","Média","Baixa"],
 "CLASSE":["A — Prioridade máxima","B — Alta prioridade","C — Prioridade de documentação","D — Monitoramento","E — Registro histórico"],
 "ABCDE":["A","B","C","D","E"],
 "SITMON":["Sem alterações","Alterações positivas","Alterações negativas","Intervenção","Degradação","Abandono","Mudança de uso","Demolição parcial","Demolição total","Risco iminente"],
 "ACAO":["Nenhuma","Nova documentação","Fiscalização","Estudo técnico","Conservação","Restauração","Proteção emergencial","Solicitação de proteção patrimonial"],
 "VIS":["Sim","Não","Não foi possível"],
 "TCLE":["Sim — termo assinado","Sim — consentimento gravado","Não obtido"],
 "FOTOH":["Existentes","Não localizadas"],
 "DANOMAP":["Sim","Não"],
}
DV={}
col=1
for name,vals in LISTS.items():
    wl.cell(1,col,name).font=F(9,True)
    for i,v in enumerate(vals): wl.cell(2+i,col,v).font=F(9)
    L=get_column_letter(col)
    DV[name]=f"Listas!${L}$2:${L}${1+len(vals)}"
    col+=1
wl.sheet_state="hidden"
dv_objs={}
def add_dv(sheet,name,cellref):
    key=(sheet.title,name)
    if key not in dv_objs:
        d=DataValidation(type="list",formula1="="+DV[name],allow_blank=True,showErrorMessage=True,
                         errorTitle="Valor inválido",error="Escolha uma opção da lista.")
        sheet.add_data_validation(d); dv_objs[key]=d
    dv_objs[key].add(cellref)

# ---------------- FIAMS ----------------
widths={"A":7,"B":38,"C":13,"D":46,"E":30,"F":44,"G":4}
for k,v in widths.items(): ws.column_dimensions[k].width=v
ws.column_dimensions["G"].hidden=True
ws.sheet_view.showGridLines=False

ws.merge_cells("A1:F1"); ws["A1"]="FIAMS — FICHA DE INVENTÁRIO DA ARQUITETURA MODERNA DE SÃO LUÍS (1930–1980)"
ws["A1"].font=F(14,True,"FFFFFF"); ws["A1"].fill=fill(C_TITLE); ws["A1"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[1].height=30
ws.merge_cells("A2:F2"); ws["A2"]="Instrumento de identificação, documentação, avaliação patrimonial e monitoramento  ·  Monte Castelo, João Paulo e Filipinho  ·  Nível 1 — Ficha de inventário"
ws["A2"].font=F(9,False,"FFFFFF",True); ws["A2"].fill=fill(C_SEC); ws["A2"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[2].height=18
# legenda
ws["A3"]="Legenda"; ws["A3"].font=F(8,True)
leg=[("B3","CAMPO — coletar na visita in loco",TAGFILL[CAMPO]),("C3","",None),("D3","DOCUMENTAL — arquivo, bibliografia, cartografia",TAGFILL[DOC]),("E3","ANÁLISE — avaliação técnica posterior",TAGFILL[ANA])]
for ref,t,c in leg:
    if c: ws[ref]=t; ws[ref].font=F(8,True,"FFFFFF"); ws[ref].fill=fill(c); ws[ref].alignment=CEN
ws.merge_cells("B3:C3")
ws["F3"]="Células em laranja-claro = registro obrigatório na visita de campo"; ws["F3"].font=F(8,True,"8A4B00"); ws["F3"].fill=fill(AMBER); ws["F3"].alignment=CEN
ws.row_dimensions[3].height=22
ws.row_dimensions[4].height=20
hdr=["Nº","Campo","Coleta","Registro / resposta","Fonte · grau de certeza","Orientação de preenchimento"]
for i,h in enumerate(hdr,1):
    c=ws.cell(5,i,h); c.font=F(9,True,"FFFFFF"); c.fill=fill(C_TITLE); c.alignment=CEN; c.border=BOX
ws.row_dimensions[5].height=20
ws.freeze_panes="A6"

r=6
KEY={}
FIELD_ITEMS=[]   # (section, code, label, row)
cur_sec=[""]; cur_sub=[""]

def section(t):
    global r
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    c=ws.cell(r,1,t); c.font=F(11,True,"FFFFFF"); c.fill=fill(C_SEC); c.alignment=Alignment(vertical="center",indent=1)
    ws.row_dimensions[r].height=22; cur_sec[0]=t; cur_sub[0]=""; r+=1
def sub(t):
    global r
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    c=ws.cell(r,1,t); c.font=F(9.5,True,C_TITLE); c.fill=fill(C_SUB); c.alignment=Alignment(vertical="center",indent=1)
    ws.row_dimensions[r].height=18; cur_sub[0]=t; r+=1
def note(t,h=30):
    global r
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    c=ws.cell(r,1,t); c.font=F(8.5,False,"4A4A4A",True); c.fill=fill(C_NOTE); c.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
    ws.row_dimensions[r].height=h; r+=1
def tagcell(row,tag):
    c=ws.cell(row,3,tag); c.font=F(7.5,True,"FFFFFF"); c.fill=fill(TAGFILL[tag]); c.alignment=CEN; c.border=BOX
def style_row(row,tag,count=True):
    campo = "CAMPO" in tag
    for cc in (1,2):
        c=ws.cell(row,cc); c.border=BOX
        if campo: c.fill=fill(AMBER_L)
    if campo: ws.cell(row,1).border=Border(left=thick_or,right=thin,top=thin,bottom=thin)
    tagcell(row,tag)
    d=ws.cell(row,4); d.border=BOX; d.alignment=WRAP; d.font=F(9,False,"000000")
    if campo: d.fill=fill(AMBER)
    e=ws.cell(row,5); e.border=BOX; e.alignment=WRAP; e.font=F(8.5)
    f=ws.cell(row,6); f.border=BOX; f.alignment=WRAP; f.font=F(8,False,"5A5A5A",True)
    ws.cell(row,7,1 if count else 0).font=F(7,c="FFFFFF")
def field(code,label,tag,dv=None,hint="",key=None,h=None,val=None,fmt=None,count=True,roteiro_label=None):
    global r
    ws.cell(r,1,code).font=F(8,False,"4A4A4A"); ws.cell(r,1).alignment=CEN
    ws.cell(r,2,label).font=F(9,True); ws.cell(r,2).alignment=WRAP
    style_row(r,tag,count)
    if hint: ws.cell(r,6,hint)
    if dv: add_dv(ws,dv,f"D{r}")
    if val is not None: ws.cell(r,4,val)
    if fmt: ws.cell(r,4).number_format=fmt
    ws.row_dimensions[r].height=h or (30 if len(hint)>60 or len(label)>45 else 18)
    if key: KEY[key]=f"D{r}"
    if "CAMPO" in tag: FIELD_ITEMS.append((cur_sec[0],code,roteiro_label or label,r))
    r+=1
def text(code,label,tag,hint="",h=60,key=None):
    field(code,label,tag,hint=hint,h=h,key=key)
    ws.cell(r-1,4).alignment=WRAPT
def checklist(code,label,tag,opts,dv="SNV",hint="Selecione para cada item."):
    global r
    ws.cell(r,1,code).font=F(8,False,"4A4A4A"); ws.cell(r,1).alignment=CEN
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6)
    c=ws.cell(r,2,f"{label}   —   {hint}"); c.font=F(9,True,C_TITLE); c.alignment=Alignment(vertical="center")
    ws.row_dimensions[r].height=17; r+=1
    for o in opts:
        ws.cell(r,2,"   ☐  "+o).font=F(9); ws.cell(r,2).alignment=WRAP
        style_row(r,tag); add_dv(ws,dv,f"D{r}")
        ws.row_dimensions[r].height=16
        if "CAMPO" in tag: FIELD_ITEMS.append((cur_sec[0],code,f"{label}: {o}",r))
        r+=1
def table(headers,rows,tags,dvs=None,count=True,rowh=18,first_label_bold=False):
    """headers: 5 strings for B,C(Coleta),D,E,F; rows: list of labels (B); dvs: dict col->list"""
    global r
    ws.cell(r,1,"").fill=fill(C_TH)
    for i,h in enumerate(headers):
        c=ws.cell(r,2+i,h); c.font=F(8.5,True,C_TITLE); c.fill=fill(C_TH); c.alignment=CEN; c.border=BOX
    ws.cell(r,1).border=BOX; ws.row_dimensions[r].height=18; r+=1
    out=[]
    for i,lab in enumerate(rows):
        tag=tags[i] if isinstance(tags,list) else tags
        ws.cell(r,2,lab).font=F(9,first_label_bold); ws.cell(r,2).alignment=WRAP
        cnt = count if not isinstance(count,list) else count[i]
        style_row(r,tag,cnt)
        ws.cell(r,6).font=F(9); 
        if "CAMPO" in tag:
            for cc in (5,6): ws.cell(r,cc).fill=fill(AMBER)
        for colL,lst in (dvs or {}).items(): add_dv(ws,lst,f"{colL}{r}")
        ws.row_dimensions[r].height=rowh
        if "CAMPO" in tag and lab and cnt: FIELD_ITEMS.append((cur_sec[0],"",f"{cur_sub[0].split(' ',1)[-1]} — {lab}",r))
        out.append(r); r+=1
    return out


SC=[]; _CTX={'in_text':False}
_o_section,_o_sub,_o_note,_o_field,_o_text,_o_check,_o_table=section,sub,note,field,text,checklist,table
def section(t): SC.append({'t':'section','title':t,'row':r}); return _o_section(t)
def sub(t): SC.append({'t':'sub','title':t,'row':r}); return _o_sub(t)
def note(t,h=30): SC.append({'t':'note','text':t,'row':r}); return _o_note(t,h)
def field(code,label,tag,dv=None,hint="",key=None,h=None,val=None,fmt=None,count=True,roteiro_label=None):
    kind='textarea' if _CTX['in_text'] else ('date' if fmt and 'YY' in fmt else ('number' if fmt else 'text'))
    SC.append({'t':'field','code':code,'label':label,'tag':tag,'dv':dv,'hint':hint,'fmt':fmt,'val':val,'count':count,'row':r,'kind':kind})
    return _o_field(code,label,tag,dv=dv,hint=hint,key=key,h=h,val=val,fmt=fmt,count=count,roteiro_label=roteiro_label)
def text(code,label,tag,hint="",h=60,key=None):
    _CTX['in_text']=True
    try: return _o_text(code,label,tag,hint=hint,h=h,key=key)
    finally: _CTX['in_text']=False
def checklist(code,label,tag,opts,dv="SNV",hint="Selecione para cada item."):
    SC.append({'t':'check','code':code,'label':label,'tag':tag,'dv':dv,'hint':hint,'options':opts,'row':r})
    return _o_check(code,label,tag,opts,dv=dv,hint=hint)
def table(headers,rows,tags,dvs=None,count=True,rowh=18,first_label_bold=False):
    SC.append({'t':'table','headers':headers,'rows':rows,'tags':tags if isinstance(tags,list) else [tags]*len(rows),'dvs':dvs or {},'count':count if isinstance(count,list) else [count]*len(rows),'row':r})
    return _o_table(headers,rows,tags,dvs=dvs,count=count,rowh=rowh,first_label_bold=first_label_bold)

# ===== 1 =====
section("1. IDENTIFICAÇÃO DO BEM")
sub("1.1 Código de inventário")
field("1.1.1","Código FIAMS",DOC,hint="Padrão: FIAMS-[bairro]-[nº sequencial]. Ex.: FIAMS-MC-001 (MC = Monte Castelo; JP = João Paulo; FL = Filipinho).",key="ID")
field("1.1.2","Nº no Inventário (planilha geral do projeto)",DOC,hint="Número/registro correspondente na planilha consolidada do inventário, para rastreabilidade cruzada.")
field("1.1.3","Código cadastral imobiliário",DOC,hint="Inscrição imobiliária municipal (IPTU), quando obtida.")
field("1.1.4","Data do primeiro registro",DOC,fmt="DD/MM/YYYY",hint="DD/MM/AAAA")
field("1.1.5","Data da última atualização",DOC,fmt="DD/MM/YYYY",hint="DD/MM/AAAA")
field("1.1.6","Versão da ficha",DOC,hint="Ex.: v1.0 (primeiro registro); v1.1 (revisão pós-campo).")
sub("1.2 Denominações")
field("1.2.1","Nome atual",CAMPO,hint="Nome em uso no local (placa, fachada, letreiro, ocupante). Ex.: Clínica X.",key="NOME_ATUAL",roteiro_label="Nome atual (placa, letreiro, ocupante)")
field("1.2.2","Nome original",DOC,key="NOME_ORIG",hint="Denominação no projeto ou na inauguração. Ex.: Residência Armando Castro.")
field("1.2.3","Nome popular",CAMPO,hint="Como o imóvel é conhecido pela vizinhança — confirmar em conversa com moradores.")
field("1.2.4","Outras denominações",DOC,key="OUTRAS")
field("1.2.5","Denominação histórica",DOC)
sub("1.3 Localização")
field("1.3.1","Logradouro",CD,key="LOGR",hint="Tipo + nome, sem abreviações. Ex.: Avenida Getúlio Vargas. Confirmar in loco.")
field("1.3.2","Número",CD,key="NUM",hint="Número afixado no imóvel. Registrar divergências com cadastro/bibliografia na coluna E.")
field("1.3.3","Complemento",CAMPO)
field("1.3.4","Bairro",CD,dv="BAIRRO",key="BAIRRO")
field("1.3.5","Localidade / subárea",CD,dv="LOCAL",hint="Subárea do recorte (Retiro Natal, Fátima, Vila Passos, Belira, Canto da Fabril, Apeadouro, Alemanha, Coreia, Vila Ivar Saldanha).")
field("1.3.6","Distrito",DOC)
field("1.3.7","CEP",DOC,hint="00000-000")
field("1.3.8","Município / UF / País",DOC,val="São Luís / Maranhão / Brasil")
field("1.3.9","Fonte do endereço",DOC,hint="Ex.: bibliografia (autor, ano), cadastro municipal, base cartográfica, verificação in loco.")
field("1.3.10","Confiança do endereço",CD,dv="CONF",hint="Alta = confirmado em campo; Média = fonte documental convergente; Baixa = fontes divergentes ou não verificado.")
sub("1.4 Georreferenciamento")
field("1.4.1","Latitude (graus decimais)",CAMPO,key="LAT",fmt="0.000000",hint="Tomar no acesso principal do lote. Ex.: -2.543210")
field("1.4.2","Longitude (graus decimais)",CAMPO,key="LON",fmt="0.000000",hint="Ex.: -44.276540")
field("1.4.3","Coordenada UTM E",CAMPO,fmt="0.00")
field("1.4.4","Coordenada UTM N",CAMPO,fmt="0.00")
field("1.4.5","Zona UTM",DOC,val="23M")
field("1.4.6","Sistema de referência",DOC,val="SIRGAS 2000")
field("1.4.7","Altitude (m)",CAMPO,fmt="0.0")
field("1.4.8","Precisão do posicionamento",CAMPO,dv="PREC")
sub("1.5 Situação patrimonial")
field("1.5.1","Proteção patrimonial (situação principal)",DOC,dv="PROT",key="PROT")
checklist("1.5.2","Outras formas de reconhecimento",DOC,["DOCOMOMO","Inventário acadêmico","Inventário comunitário","Publicação especializada","Reconhecimento por instituição","Outro (especificar na coluna E)"],dv="SN")
field("1.5.3","Instrumento / processo",DOC)
field("1.5.4","Número",DOC)
field("1.5.5","Órgão",DOC,hint="Ex.: IPHAN, DPHAP-MA, FUMPH (municipal).")
field("1.5.6","Ano",DOC)

# ===== 2 =====
section("2. CARACTERIZAÇÃO DO BEM")
sub("2.1 Natureza")
field("2.1.1","Natureza do bem",CAMPO,dv="NAT")
field("2.1.2","Categoria funcional (original)",DOC,dv="CAT")
field("2.1.3","Categoria funcional (atual)",CAMPO,dv="CAT")
sub("2.2 Tipologia arquitetônica")
field("2.2.1","Tipologia original",DOC,key="TIPO",hint="Ex.: residência unifamiliar térrea; edifício escolar pavilhonar.")
field("2.2.2","Tipologia atual",CAMPO)
field("2.2.3","Subtipologia",ANA)
sub("2.3 Relação com a produção moderna")
field("2.3.1","Classificação principal",ANA,dv="LING",key="LING")
field("2.3.2","Classificação secundária (se houver)",ANA,dv="LING")
text("2.3.3","Justificativa da classificação",ANA,hint="Relacionar a classificação a atributos observados em campo (seção 8) e a fontes.")
note("Importante: a classificação estilística não deve ser utilizada isoladamente para determinar a inclusão do imóvel. A ficha deve registrar também aspectos históricos, sociais, urbanos, tecnológicos e documentais.",24)

# ===== 3 =====
section("3. AUTORIA E AGENTES DA PRODUÇÃO")
note("A autoria desconhecida NÃO é critério de exclusão. O inventário admite o moderno difuso das periferias, cuja produção frequentemente não tem autoria ou data documentadas.",24)
sub("3.1 Projeto arquitetônico")
field("3.1.1","Arquiteto(a)",DOC,key="ARQ",hint="Nome completo. Se desconhecido, registrar 'Não identificado'.")
field("3.1.2","Registro profissional",DOC)
field("3.1.3","Escritório / equipe",DOC)
field("3.1.4","Autoria confirmada?",DOC,dv="AUT",hint="Sim = fonte primária; Provável = fonte secundária; Atribuição = análise formal/comparativa.")
field("3.1.5","Fonte da autoria",DOC)
field("3.1.6","Placa, assinatura ou marca de autoria visível no imóvel",CAMPO,hint="Verificar placas de obra, inscrições em pisos, painéis ou fachadas.")
sub("3.2 Outros profissionais")
for i,l in enumerate(["Engenheiro estrutural","Engenheiro de instalações","Paisagista","Artista plástico","Designer / mobiliário","Construtor / empreiteiro","Fotógrafo / documentarista"],1):
    field(f"3.2.{i}",l,DOC)
sub("3.3 Contratante")
field("3.3.1","Proprietário original",DOC)
field("3.3.2","Instituição / empresa",DOC)
field("3.3.3","Agente promotor",DOC)
field("3.3.4","Natureza do contratante",DOC,dv="NATC")
sub("3.4 Autoria intelectual, relações e difusão")
text("3.4.1","Relação do edifício com a produção do autor / escritório",ANA,h=45)
text("3.4.2","Obras relacionadas do mesmo autor",DOC,h=40)
field("3.4.3","Vetor de difusão do moderno",ANA,dv="VETOR",hint="Canal pelo qual a linguagem moderna chegou ao bem.")

# ===== 4 =====
section("4. CRONOLOGIA DO BEM")
rows=["Projeto","Aprovação","Início da construção","Conclusão","Inauguração","Primeira alteração","Segunda alteração","Outras intervenções","Mudança de uso","Restauro / reforma","Situação atual"]
tg=[DOC]*5+[CD]*5+[CAMPO]
crono=table(["Evento","Coleta","Data","Fonte","Grau de certeza"],rows,tg,dvs={"F":"CERT"})
KEY["ANO_PROJ"]=f"D{crono[0]}"; KEY["ANO_CONC"]=f"D{crono[3]}"
field("4.1","Data provável de construção",CD,hint="Admite aproximação: c. 1965; década de 1960. Explicitar o critério na coluna E.")
field("4.2","Período",DOC,dv="PER")

# ===== 5 =====
section("5. HISTÓRIA E DESENVOLVIMENTO")
sub("5.1 História da edificação")
note("Registrar de forma crítica e documentada: origem, motivação da construção, contratante, projeto, construção, inauguração, ocupações, proprietários, usos, transformações e acontecimentos relevantes. Divergências entre fontes devem ser explicitadas, não resolvidas silenciosamente.",30)
text("5.1.1","Síntese histórica",DOC,h=110)
sub("5.2 História urbana")
text("5.2.1","Relação com o processo de modernização de São Luís",ANA,h=50)
text("5.2.2","Relação com o bairro",ANA,h=40)
text("5.2.3","Relação com eixos de expansão urbana",ANA,h=40,hint="Ex.: eixo Caminho Grande / Av. Getúlio Vargas – Av. João Pessoa.")
text("5.2.4","Relação com políticas públicas, planos urbanos ou processos econômicos",ANA,h=45)

# ===== 6 =====
section("6. DOCUMENTAÇÃO HISTÓRICA")
checklist("6.1","Fontes primárias localizadas",DOC,["Projeto arquitetônico","Plantas originais","Fachadas","Cortes","Memorial descritivo","Fotografias de época","Fotografias de construção","Licença / alvará","Registro imobiliário","Jornal","Revista","Correspondência","Mapas","Cadastro municipal","Arquivo do proprietário","Arquivo do arquiteto","Depoimento oral","Outro (especificar na coluna E)"],dv="FONTE",hint="Localizada / Não localizada / Não pesquisada.")
sub("6.2 Localização das fontes")
for i,l in enumerate(["Instituição / arquivo","Fundo / coleção","Documento","Código / referência"],1): field(f"6.2.{i}",l,DOC)
field("6.2.5","Data",DOC)
sub("6.3 Fotografias históricas")
field("6.3.1","Quantidade",DOC,fmt="0")
field("6.3.2","Período mais antigo documentado",DOC)
field("6.3.3","Fonte",DOC)
field("6.3.4","Série temporal digital (Street View / Google Earth)",DOC,hint="Anos disponíveis. Ex.: 2011, 2014, 2019, 2024.")
field("6.3.5","Existem imagens que permitam comparação temporal?",DOC,dv="COMP")
field("6.3.6","Fotografias antigas em posse de moradores / proprietários",CAMPO,dv="SNV",hint="Perguntar durante a visita; se houver, solicitar autorização para reprodução.")
note("Guedes identifica a ausência de fotografias históricas e desenhos originais como deficiência recorrente nos inventários brasileiros, dificultando a comparação das diferentes fases do edifício.",24)

# ===== 7 =====
section("7. IMPLANTAÇÃO E RELAÇÃO COM O LOTE")
sub("7.1 Características do lote")
field("7.1.1","Área do lote (m²)",CD,key="AREA_LOTE",fmt="#,##0.00",hint="Cadastro municipal ou medição/estimativa em campo; indicar o método na coluna E.")
field("7.1.2","Testada (m)",CAMPO,fmt="0.00",hint="Medir com trena a laser no alinhamento.")
field("7.1.3","Profundidade (m)",CD,fmt="0.00")
field("7.1.4","Forma do lote",CAMPO,dv="FORMA")
sub("7.2 Implantação")
field("7.2.1","Tipo de implantação",CAMPO,dv="IMPL")
field("7.2.2","Recuo frontal (m)",CAMPO,fmt="0.00")
field("7.2.3","Recuo lateral (m)",CAMPO,fmt="0.00",hint="Se diferentes, registrar esquerdo/direito na coluna E.")
field("7.2.4","Recuo posterior (m)",CAMPO,fmt="0.00")
sub("7.3 Relação edifício–lote–rua")
for i,l in enumerate(["Implantação","Acessos (pedestre / veículo)","Recuos (uso atual)","Jardins","Áreas livres","Muros (altura, material, data provável)","Cercamentos (grades, cerca elétrica, gradis)","Relação com a calçada","Relação com o alinhamento","Relação visual com a rua"],1):
    field(f"7.3.{i}",l,CAMPO,h=22)
text("7.3.11","Análise",ANA,h=60)

# ===== 8 =====
section("8. CARACTERIZAÇÃO ARQUITETÔNICA")
sub("8.1 Volumetria")
field("8.1.1","Número de pavimentos",CAMPO,key="PAV",fmt="0")
field("8.1.2","Altura aproximada (m)",CAMPO,fmt="0.0")
field("8.1.3","Área construída aproximada (m²)",CD,key="AREA_CONST",fmt="#,##0.00")
for i,l in enumerate(["Subsolo","Pilotis","Mezanino","Terraço"],4): field(f"8.1.{i}",l,CAMPO,dv="SNV")
text("8.1.8","Descrição volumétrica",CAMPO,h=50)
sub("8.2 Organização espacial")
field("8.2.1","Acesso ao interior durante a visita",CAMPO,dv="ACESSO",hint="Condiciona os campos 8.2, 9.5 e 13.3. Registrar o motivo em caso de acesso parcial ou negado.")
text("8.2.2","Partido arquitetônico",CD,h=45)
text("8.2.3","Circulação",CAMPO,h=35)
text("8.2.4","Setorização funcional",CAMPO,h=35)
text("8.2.5","Relação entre espaços internos e externos",CAMPO,h=35)
checklist("8.3","Características formais",CAMPO,["Horizontalidade","Verticalidade","Composição assimétrica","Planta livre","Estrutura independente","Pilotis","Laje plana","Cobertura-terraço","Telhado borboleta","Platibanda","Brise-soleil","Elementos vazados / cobogós","Esquadrias moduladas","Panos de vidro","Marquise","Balanço","Rampa","Escada escultórica","Integração interior–exterior","Jardim integrado","Outras características (coluna E)"])
text("8.3.1","Descrição analítica",ANA,h=60)

# ===== 9 =====
section("9. CARACTERIZAÇÃO TÉCNICO-CONSTRUTIVA")
sub("9.1 Sistema estrutural")
field("9.1.1","Sistema estrutural predominante",CAMPO,dv="ESTR")
text("9.1.2","Sistema estrutural identificado (descrição)",CAMPO,h=35,hint="Pilares, vigas, lajes aparentes; vãos; juntas; indícios observáveis.")
sub("9.2 Vedações")
text("9.2.1","Vedações",CAMPO,h=35,hint="Material, espessura aparente, alvenarias originais ou substituídas.")
sub("9.3 Cobertura")
text("9.3.1","Cobertura",CAMPO,h=35,hint="Ex.: laje plana impermeabilizada; telhado oculto por platibanda; telhado borboleta; telha de fibrocimento.")
sub("9.4 Esquadrias")
field("9.4.1","Material",CAMPO,hint="Ex.: madeira, ferro, alumínio, vidro.")
field("9.4.2","Sistema",CAMPO,hint="Ex.: basculante, pivotante, veneziana de correr, fixa.")
field("9.4.3","Estado de integridade",CAMPO,dv="ESQ")
sub("9.5 Revestimentos")
for i,l in enumerate(["Fachadas","Pisos","Paredes internas","Forros"],1): field(f"9.5.{i}",l,CAMPO,h=22)
checklist("9.6","Elementos de controle ambiental",CAMPO,["Brises","Cobogós","Varandas","Venezianas","Pátios","Ventilação cruzada","Proteção solar","Beirais","Outros (coluna E)"])
text("9.6.1","Desempenho climático observado",CD,h=45,hint="Orientação solar das fachadas, sombreamento, ventilação percebida, adaptações posteriores (ar-condicionado, películas, fechamentos).")
note("Este campo recebe atenção especial em São Luís: Guedes aponta que muitos inventários registram a forma, mas pouco exploram a relação da obra com os condicionantes climáticos e legais.",24)

# ===== 10 =====
section("10. BENS INTEGRADOS E SÍNTESE DAS ARTES")
note("Registrar elementos concebidos como parte integrante do projeto. Os bens integrados constituem camada patrimonial autônoma, e não mero atributo subordinado da edificação.",24)
checklist("10.1","Bens integrados identificados",CAMPO,["Painéis","Azulejaria","Mosaicos","Murais","Esculturas","Obras de arte","Vitral","Mobiliário fixo","Luminárias originais","Paisagismo","Comunicação visual","Elementos gráficos","Outro (coluna E)"])
text("10.2","Descrição",CAMPO,h=50,hint="Localização no edifício, dimensões aproximadas, técnica, cores, motivos.")
field("10.3","Autoria",DOC)
field("10.4","Estado de conservação",CAMPO,dv="CONS")
field("10.5","Código no Inventário da Azulejaria Modernista (se houver)",DOC)

# ===== 11 =====
section("11. USO E APROPRIAÇÃO")
field("11.1","Uso original",DOC,key="USO_ORIG")
field("11.2","Uso atual",CAMPO,key="USO_ATUAL",hint="Atividade efetivamente observada no dia da visita.")
field("11.2.1","Ocupante / usuário atual",CAMPO)
sub("11.3 Histórico de usos")
table(["Período","Coleta","Uso","Responsável / ocupante","Fonte"],["","","",""],CD,count=False)
sub("11.4 Compatibilidade do uso atual")
field("11.4.1","Compatibilidade",ANA,dv="COMPAT")
text("11.4.2","Impactos do uso atual sobre o edifício",CD,h=45,hint="Ex.: fechamento de varandas, letreiros sobre a fachada, ocupação do recuo por estacionamento.")

# ===== 12 =====
section("12. RELAÇÃO COM O ENTORNO E A PAISAGEM URBANA")
note("Seção mais desenvolvida que nas fichas tradicionais do DOCOMOMO, pois Guedes identifica a ambiência e a relação com o entorno como lacunas recorrentes.",20)
sub("12.1 Caracterização do entorno imediato")
for i,(l,h) in enumerate([("Uso predominante","Ex.: comércio e serviços no térreo; residencial."),("Gabarito predominante","Nº de pavimentos predominante na quadra e nas testadas opostas."),("Densidade",""),("Morfologia","Ex.: lotes estreitos entre medianeiras; lotes amplos com recuos."),("Estado de conservação predominante","")],1):
    field(f"12.1.{i}",l,CAMPO,hint=h)
sub("12.2 Relações urbanas")
for i,l in enumerate(["Relação com a quadra","Relação com a rua","Relação com edificações vizinhas","Relação visual","Relação funcional"],1):
    text(f"12.2.{i}",l,CAMPO,h=30)
text("12.2.6","Contribuição para a paisagem urbana",ANA,h=40)
checklist("12.3","Ambiência",CAMPO,["Elemento dominante","Elemento focal","Elemento integrante de conjunto","Elemento de referência urbana","Elemento isolado","Elemento descaracterizado pelo entorno"],dv="SN")

# ===== 13 =====
section("13. ANÁLISE DA AUTENTICIDADE")
for i,l in enumerate(["Autenticidade da implantação","Autenticidade volumétrica","Autenticidade espacial","Autenticidade dos materiais","Autenticidade dos elementos arquitetônicos","Autenticidade do uso"],1):
    field(f"13.{i}",l,ANA,dv="AUTENT")
text("13.7","Síntese",ANA,h=50)

# ===== 14 =====
section("14. INTEGRIDADE")
sub("14.1 Integridade por dimensão")
for i,l in enumerate(["Integridade do lote","Integridade da implantação","Integridade volumétrica","Integridade espacial","Integridade construtiva","Integridade material","Integridade dos elementos caracterizadores","Integridade paisagística","Integridade da relação com o entorno"],1):
    field(f"14.1.{i}",l,ANA,dv="INTEG")
field("14.1.10","Grau geral de integridade",ANA,dv="GRAU",key="INTEG")
text("14.1.11","Justificativa",ANA,h=45)
sub("14.2 Índice de Integridade Modernista (IIM) — atributos A1–A7")
note("Estado de cada atributo: Preservado = 2; Adaptado = 1; Suprimido / Ausente = 0. 'n.a.' (atributo inexistente no projeto original) e 's.d.' (atributo não descrito / não verificável) são excluídos do cálculo. IIM = pontos obtidos ÷ (nº de atributos aplicáveis × 2).",30)
attrs=["A1 — Composição volumétrica e geometria","A2 — Cobertura plana / platibanda","A3 — Aberturas modernistas","A4 — Sistema estrutural e organização da planta","A5 — Materiais em estado natural e elementos vazados","A6 — Ausência de ornamentação","A7 — Integração com paisagismo e sítio"]
iim_rows=table(["Atributo","Coleta","Estado do atributo","Pontos","Evidência observada"],attrs,CAMPO,dvs={"D":"IIM"},rowh=20)
for rr in iim_rows:
    ws.cell(rr,5,f'=IF(D{rr}="Preservado",2,IF(D{rr}="Adaptado",1,IF(D{rr}="Suprimido / Ausente",0,"")))')
    ws.cell(rr,5).alignment=CEN; ws.cell(rr,5).fill=fill("FFFFFF"); ws.cell(rr,5).font=F(9,True)
a,b=iim_rows[0],iim_rows[-1]
nexpr=f'(COUNTIF(D{a}:D{b},"Preservado")+COUNTIF(D{a}:D{b},"Adaptado")+COUNTIF(D{a}:D{b},"Suprimido / Ausente"))'
field("14.2.1","Atributos aplicáveis (nº)",ANA)
ws.cell(r-1,4,f"={nexpr}"); ws.cell(r-1,6,"Calculado automaticamente.")
field("14.2.2","IIM (%)",ANA,fmt="0.0%",key="IIM")
ws.cell(r-1,4,f'=IF({nexpr}=0,"",SUM(E{a}:E{b})/({nexpr}*2))'); ws.cell(r-1,6,"Calculado automaticamente.")
iimref=f"D{r-1}"
field("14.2.3","Categoria IIM",ANA,key="IIM_CAT")
ws.cell(r-1,4,f'=IF({iimref}="","",IF({iimref}>=0.85,"Íntegro (muito preservado)",IF({iimref}>=0.7,"Preservado",IF({iimref}>=0.55,"Preservado com adaptações",IF({iimref}>=0.35,"Parcialmente preservado",IF({iimref}>=0.15,"Descaracterizado","Não preservado / substituído"))))))')
ws.cell(r-1,6,"Faixas: 85–100% Íntegro; 70–84% Preservado; 55–69% Preservado com adaptações; 35–54% Parcialmente preservado; 15–34% Descaracterizado; 0–14% Não preservado. Bens demolidos: registrar 'Demolido' na coluna E.")
ws.row_dimensions[r-1].height=40
for rr in (r-3,r-2,r-1): ws.cell(rr,4).font=F(9,True,C_TITLE)

# ===== 15 =====
section("15. ALTERAÇÕES E TRANSFORMAÇÕES")
sub("15.1 Alterações entre projeto e obra")
table(["Elemento","Coleta","Projeto","Obra executada","Data · Fonte"],["Implantação","Volumetria","Fachada","Planta","Materiais","Estrutura"],CD)
sub("15.2 Alterações posteriores")
note("Grau: Conservação, Manutenção, Reforma, Ampliação, Adaptação, Acréscimo, Supressão, Substituição, Demolição parcial.  Reversibilidade: Reversível, Parcialmente reversível, Irreversível.",20)
alt=table(["Ano — intervenção","Coleta","Elemento afetado","Grau da intervenção","Reversibilidade"],["","","","","",""],CD,dvs={"E":"INTERV","F":"REV"},count=[True]+[False]*5,rowh=22)
FIELD_ITEMS.append(("15. ALTERAÇÕES E TRANSFORMAÇÕES","15.2","Alterações posteriores visíveis (acréscimos, supressões, substituições) — registrar cada uma",alt[0]))
note("Guedes considera a identificação das alterações significativas, com suas respectivas datas, um dos diferenciais das fichas DOCOMOMO em relação ao SICG/IPHAN.",20)

# ===== 16 =====
section("16. ESTADO DE CONSERVAÇÃO")
field("16.1","Avaliação geral",CAMPO,dv="CONS",key="CONS")
sub("16.2 Avaliação por componente")
table(["Componente","Coleta","Estado","Patologias observadas","Observações"],["Estrutura","Cobertura","Fachadas","Esquadrias","Revestimentos","Pisos","Instalações","Elementos integrados","Paisagismo"],CAMPO,dvs={"D":"COMPO"})

# ===== 17 =====
section("17. MAPA DE DANOS E PATOLOGIAS")
table(["Agente identificado","Coleta","Presença","Localização / componente","Intensidade"],["Umidade","Infiltração","Eflorescência","Fissuração","Trincas","Corrosão","Desagregação","Perda de revestimento","Ataque biológico","Vegetação invasiva","Poluição","Intervenção inadequada","Sobrecarga","Falha de manutenção","Outro"],CAMPO,dvs={"D":"PATO","F":"INT3"},rowh=16)
field("17.1","Mapa de danos anexado",DOC,dv="DANOMAP")
field("17.2","Arquivo SIG/GIS",DOC,hint="Caminho ou nome do arquivo (.gpkg, .shp, .qgz).")

# ===== 18 =====
section("18. RISCOS E AMEAÇAS")
table(["Ameaça","Coleta","Intensidade","Urgência","Tendência"],["Abandono","Demolição","Reforma descaracterizante","Alteração de uso","Perda de elementos","Pressão imobiliária","Adensamento do entorno","Problemas estruturais","Falta de manutenção","Risco climático","Outros"],[CAMPO]*11,dvs={"D":"INT3","E":"INT3","F":"TEND"},rowh=16)
field("18.1","Classificação de risco",ANA,dv="RISCO",key="RISCO")

# ===== 19 =====
section("19. VALORES CULTURAIS E SIGNIFICÂNCIA")
note("Núcleo analítico da ficha. Para cada valor, atribuir o grau e registrar a justificativa na linha seguinte.",18)
vals=[("19.1","Valor histórico","V_HIST"),("19.2","Valor arquitetônico","V_ARQ"),("19.3","Valor urbanístico","V_URB"),("19.4","Valor tecnológico / construtivo",None),("19.5","Valor paisagístico",None),("19.6","Valor social","V_SOC"),("19.7","Valor simbólico / memorial",None),("19.8","Valor documental","V_DOC")]
for code,l,k in vals:
    field(code,l,ANA,dv="VAL",key=k)
    text(code+".j","    Justificativa",ANA,h=36)
field("19.9","Valor de representatividade do modernismo maranhense",ANA,dv="VAL5",key="SIGN")
text("19.9.j","    Justificativa",ANA,h=36)
text("19.10","Declaração de significância",ANA,h=110,hint="Por que este edifício é importante para a história da arquitetura e da cidade de São Luís? Texto conclusivo, em 5–10 linhas.")
note("A declaração de significância responde à lacuna identificada por Guedes: valor e significância aparecem com frequência nas fichas inspiradas no DOCOMOMO, mas estão ausentes em vários instrumentos mais descritivos.",24)

# ===== 20 =====
section("20. AVALIAÇÃO DO EXEMPLAR MODERNISTA")
field("20.1","Representatividade",ANA,dv="REPR")
field("20.2","Raridade",ANA,dv="RAR")
field("20.3","Integridade",ANA,dv="AMA",hint="Coerente com as seções 14.1 e 14.2.")
field("20.4","Autenticidade",ANA,dv="AMA",key="AUTENT",hint="Coerente com a seção 13.")
field("20.5","Importância para a história urbana",ANA,dv="HU")

# ===== 21 =====
section("21. CLASSIFICAÇÃO DE PRIORIDADE PARA PRESERVAÇÃO")
note("O inventário não produz um 'valor' único, mas uma classe de prioridade, evitando hierarquia exclusivamente estética.  A — Prioridade máxima: excepcional significância, elevada integridade/autenticidade e/ou elevado risco.  B — Alta prioridade: bem representativo, com atributos relevantes e necessidade de acompanhamento.  C — Prioridade de documentação: relevante para a compreensão do conjunto, ainda que parcialmente alterado.  D — Monitoramento: baixo grau de integridade ou significância, mas cuja transformação deve ser acompanhada.  E — Registro histórico: documentação importante mesmo diante de descaracterização ou perda.",62)
field("21.1","Classe atribuída",ANA,dv="CLASSE",key="PRIOR")
text("21.2","Justificativa",ANA,h=50)

# ===== 22 =====
section("22. DOCUMENTAÇÃO GRÁFICA OBRIGATÓRIA (DOSSIÊ — NÍVEL 2)")
field("22.0","Pasta / link do dossiê gráfico",DOC,key="LINK",hint="Caminho da pasta ou URL do dossiê associado.")
checklist("22.1","Cartografia",DOC,["Mapa de localização municipal","Mapa de localização no bairro","Planta de situação","Planta cadastral","Ortofoto","Mapa de uso do solo","Mapa de gabarito","Mapa de tipologia","Mapa de proteção patrimonial","Mapa de entorno / ambiência"],dv="DOCG",hint="Sim / Não / Em elaboração / Não se aplica. Nome do arquivo na coluna E.")
checklist("22.2","Desenhos arquitetônicos",CD,["Croqui de implantação (feito em campo)","Planta de implantação","Plantas baixas","Planta de cobertura","Cortes","Fachadas","Detalhes significativos","Redesenho interpretativo"],dv="DOCG",hint="Sim / Não / Em elaboração / Não se aplica.")
note("Guedes destaca que plantas, cortes, fachadas e redesenhos constituem contribuição importante para documentação e preservação, embora estejam ausentes ou sejam insuficientes em muitos inventários analisados.",20)
checklist("22.3","Registro fotográfico atual",CAMPO,["Fachada principal","Fachadas laterais","Fundos","Cobertura","Implantação","Espaços internos","Circulação","Elementos construtivos","Elementos modernistas","Bens integrados","Entorno imediato","Paisagem urbana"],dv="VIS",hint="Registrado? Números das fotos na coluna E (ver log fotográfico na aba Roteiro_Campo).")
field("22.3.1","Fotografia principal (arquivo)",CAMPO,key="FOTO",hint="Ex.: FIAMS-MC-001_F01.jpg — fachada principal, vista frontal.")
field("22.3.2","Fotografias históricas",DOC,dv="FOTOH")

# ===== 23 =====
section("23. ENTREVISTA E MEMÓRIA SOCIAL")
note("Inspirado no INBI-SU, que incorporou entrevistas para compreender moradores, trabalhadores, usos e apropriações do sítio urbano. Obter consentimento antes de gravar ou registrar o depoimento.",22)
field("23.1","Entrevistado(a)",CAMPO)
field("23.2","Relação com o imóvel",CAMPO,hint="Ex.: proprietário, morador, ocupante, vizinho, ex-funcionário.")
field("23.3","Data",CAMPO,fmt="DD/MM/YYYY")
field("23.4","Entrevistador(a)",CAMPO)
field("23.5","Consentimento registrado",CAMPO,dv="TCLE")
sub("23.6 Questões essenciais")
qs=["Qual a relação do entrevistado com o edifício?","Quais usos o edifício já teve?","Quais transformações foram percebidas?","Existem elementos considerados particularmente importantes?","Há memórias, acontecimentos ou personagens associados ao imóvel?","Como o edifício é percebido pela comunidade?","Existem conflitos relacionados à sua conservação ou uso?"]
for i,q in enumerate(qs,1): text(f"Q{i}",q,CAMPO,h=36)
text("23.7","Síntese do depoimento",ANA,h=70)

# ===== 24 =====
section("24. FONTES E REFERÊNCIAS")
for i,l in enumerate(["Fontes arquivísticas","Fontes bibliográficas","Fontes cartográficas","Fontes iconográficas","Fontes orais","Fontes digitais","Pesquisas acadêmicas relacionadas"],1):
    text(f"24.{i}",l,DOC,h=36,hint="Formato ABNT (NBR 6023)." if i==2 else "")

# ===== 25 =====
section("25. EQUIPE E RESPONSABILIDADE TÉCNICA")
sub("25.1 Responsável pelo levantamento")
field("25.1.1","Nome",DOC); field("25.1.2","Formação",DOC); field("25.1.3","CAU / registro",DOC)
sub("25.2 Responsáveis por etapa")
for i,l in enumerate(["Pesquisa histórica","Levantamento arquitetônico (campo)","SIG / geoprocessamento","Avaliação patrimonial"],1): field(f"25.2.{i}",l,DOC)
field("25.3","Instituição",DOC,val="LUPA — Laboratório de Urbanismo, Paisagem, Arquitetura e Artes · Centro Universitário UNDB")
field("25.4","Data",DOC,fmt="DD/MM/YYYY"); field("25.5","Revisão",DOC); field("25.6","Data da revisão",DOC,fmt="DD/MM/YYYY")

# ===== 26-27 =====
section("26–27. MONITORAMENTO PERIÓDICO E MATRIZ DE MONITORAMENTO (NÍVEL 3)")
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
c=ws.cell(r,1,"→ Registros de inspeção e matriz quantitativa na aba Monitoramento (clique para abrir)."); c.hyperlink="#Monitoramento!A1"; c.font=F(9,True,"1F5FA8",u="single"); c.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20; r+=1

# ===== 28 =====
section("28. SÍNTESE FINAL DO INVENTÁRIO")
text("28.1","Atributo principal do bem",ANA,h=36)
sub("28.2 Principais atributos a preservar")
for i in range(1,6): field(f"28.2.{i}",f"Atributo {i}",ANA,h=20)
sub("28.3 Principais ameaças")
for i in range(1,4): field(f"28.3.{i}",f"Ameaça {i}",ANA,h=20)
text("28.4","Recomendações",ANA,h=60)
field("28.5","Prioridade de conservação",ANA,dv="ABCDE")
field("28.6","Prioridade de pesquisa",ANA,dv="ABCDE")
field("28.7","Prioridade de monitoramento",ANA,dv="ABCDE")
LAST=r-1
note("FIAMS — Ficha de Inventário da Arquitetura Moderna de São Luís (1930–1980). Adaptação temática e crítica dos referenciais INBI/INBI-SU e SICG/IPHAN, complementados pelas fichas DOCOMOMO e pelas experiências brasileiras de inventariação da arquitetura moderna. Não constitui reprodução de formulário oficial do IPHAN.",30)

# painel de preenchimento (linha 4)
rg=lambda L:f"${L}$6:${L}${LAST}"
def pct(tagword):
    return (f'SUMPRODUCT(({rg("G")}=1)*ISNUMBER(SEARCH("{tagword}",{rg("C")}))*({rg("D")}<>""))'
            f'/MAX(1,SUMPRODUCT(({rg("G")}=1)*ISNUMBER(SEARCH("{tagword}",{rg("C")}))))')
ws["A4"]="Painel"; ws["A4"].font=F(8,True)
ws["B4"]="Preenchimento (itens registrados / total)"; ws["B4"].font=F(8.5,True,C_TITLE); ws["B4"].alignment=WRAP
ws["C4"]=f'="Campo "&TEXT({pct("CAMPO")},"0%")'
ws["D4"]=f'="Documental "&TEXT({pct("DOC")},"0%")&"    ·    Análise "&TEXT({pct("ANÁLISE")},"0%")'
ws["E4"]=f'="Bem: "&IF({KEY["ID"]}="","—",{KEY["ID"]})'
ws["F4"]=f'="IIM: "&IF({KEY["IIM"]}="","—",TEXT({KEY["IIM"]},"0%")&" · "&{KEY["IIM_CAT"]})'
for ref,cfill,fc in (("C4",TAGFILL[CAMPO],"FFFFFF"),("D4","EAF0F4",C_TITLE),("E4","EAF0F4",C_TITLE),("F4","EAF0F4",C_TITLE)):
    ws[ref].font=F(8.5,True,fc); ws[ref].fill=fill(cfill); ws[ref].alignment=CEN; ws[ref].border=BOX

ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=9
ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.print_title_rows="5:5"; ws.print_area=f"A1:F{r}"
ws.page_margins=PageMargins(left=0.4,right=0.4,top=0.5,bottom=0.5)
ws.oddFooter.center.text="FIAMS · &P / &N"; ws.oddFooter.center.font="Arial"

# ---------------- ROTEIRO DE CAMPO ----------------
R=wr; R.sheet_view.showGridLines=False
for k,v in {"A":7,"B":58,"C":12,"D":14,"E":12,"F":40}.items(): R.column_dimensions[k].width=v
R.merge_cells("A1:F1"); R["A1"]="ROTEIRO DE CAMPO — FIAMS"; R["A1"].font=F(14,True,"FFFFFF"); R["A1"].fill=fill(TAGFILL[CAMPO]); R["A1"].alignment=Alignment(vertical="center",indent=1); R.row_dimensions[1].height=30
R.merge_cells("A2:F2"); R["A2"]="Checklist de tudo o que deve ser observado, medido, fotografado ou perguntado in loco. Registre as respostas nas células laranja da aba FIAMS (link na coluna C)."
R["A2"].font=F(9,False,"FFFFFF",True); R["A2"].fill=fill("B35F00"); R["A2"].alignment=Alignment(vertical="center",indent=1,wrap_text=True); R.row_dimensions[2].height=22
DEN=f'IF(FIAMS!{KEY["NOME_ATUAL"]}<>"",FIAMS!{KEY["NOME_ATUAL"]},IF(FIAMS!{KEY["NOME_ORIG"]}<>"",FIAMS!{KEY["NOME_ORIG"]},FIAMS!{KEY["OUTRAS"]}&""))'
hdrinfo=[("Bem (código)",f'=IF(FIAMS!{KEY["ID"]}="","",FIAMS!{KEY["ID"]})'),("Denominação","="+DEN),
 ("Endereço",f'=TRIM(FIAMS!{KEY["LOGR"]}&IF(FIAMS!{KEY["NUM"]}="","",", nº "&FIAMS!{KEY["NUM"]}))'),
 ("Data da visita",None),("Horário (início – fim)",None),("Equipe",None),("Condições climáticas / luz",None),("Autorização do proprietário / ocupante",None),("Contato no local (nome / telefone)",None)]
rr=4
for lab,fx in hdrinfo:
    R.cell(rr,1,"").fill=fill(AMBER_L)
    R.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=2)
    R.cell(rr,1,lab).font=F(9,True); R.cell(rr,1).alignment=Alignment(indent=1,vertical="center")
    R.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=6)
    c=R.cell(rr,3,fx); c.font=F(9,True if fx else False); c.fill=fill("FFFFFF" if fx else AMBER); c.border=BOX; c.alignment=WRAP
    for cc in range(1,7): R.cell(rr,cc).border=BOX
    R.row_dimensions[rr].height=18; rr+=1
R.cell(7,3).number_format="DD/MM/YYYY"
add_dv(R,"VIS","C11")
rr+=1
R.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
R.cell(rr,1,"EQUIPAMENTOS E PREPARAÇÃO").font=F(10,True,"FFFFFF"); R.cell(rr,1).fill=fill(C_SEC); R.cell(rr,1).alignment=Alignment(indent=1,vertical="center"); rr+=1
equip=["Ficha FIAMS impressa ou em tablet + prancheta","Trena a laser e trena de fita (≥ 5 m)","GPS / celular com app de coordenadas (SIRGAS 2000)","Câmera com bateria e cartão de reserva; escala métrica para detalhes","Bússola (orientação solar das fachadas)","Papel para croqui de implantação","Termo de consentimento para entrevista; gravador","Crachá / carta de apresentação institucional (LUPA/UNDB)","Mapa de localização e fotografias históricas do bem (comparação in loco)"]
for e in equip:
    R.cell(rr,2,"☐  "+e).font=F(9); R.cell(rr,4).border=BOX; add_dv(R,"SN",f"D{rr}")
    R.cell(rr,4).fill=fill(AMBER); R.cell(rr,3,"Levado?").font=F(8,False,"5A5A5A",True); rr+=1
rr+=1
# progress
R.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=2)
R.cell(rr,1,"Itens verificados no roteiro").font=F(9,True); prog_row=rr; rr+=2
# header
for i,h in enumerate(["Nº","Item a verificar in loco","Ficha","Verificado?","Foto nº","Observações de campo"],1):
    c=R.cell(rr,i,h); c.font=F(9,True,"FFFFFF"); c.fill=fill(C_TITLE); c.alignment=CEN; c.border=BOX
hdr_row=rr; rr+=1
first_item=rr; last_sec=None
for sec,code,lab,frow in FIELD_ITEMS:
    if sec!=last_sec:
        R.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
        R.cell(rr,1,sec).font=F(9.5,True,"8A4B00"); R.cell(rr,1).fill=fill(AMBER); R.cell(rr,1).alignment=Alignment(indent=1,vertical="center")
        R.row_dimensions[rr].height=18; rr+=1; last_sec=sec
    R.cell(rr,1,code).font=F(8,False,"4A4A4A"); R.cell(rr,1).alignment=CEN
    R.cell(rr,2,lab).font=F(9); R.cell(rr,2).alignment=WRAP
    c=R.cell(rr,3,"→ ficha"); c.hyperlink=f"#FIAMS!D{frow}"; c.font=F(8.5,False,"1F5FA8",u="single"); c.alignment=CEN
    add_dv(R,"VIS",f"D{rr}")
    for cc in range(1,7): R.cell(rr,cc).border=BOX
    R.cell(rr,4).fill=fill(AMBER_L)
    R.row_dimensions[rr].height=16 if len(lab)<70 else 28
    rr+=1
last_item=rr-1
R.merge_cells(start_row=prog_row,start_column=3,end_row=prog_row,end_column=4)
R.cell(prog_row,3,f'=COUNTIF(D{first_item}:D{last_item},"Sim")/MAX(1,COUNTA(B{first_item}:B{last_item})-0)')
R.cell(prog_row,3).number_format="0%"; R.cell(prog_row,3).font=F(11,True,"FFFFFF"); R.cell(prog_row,3).fill=fill(TAGFILL[CAMPO]); R.cell(prog_row,3).alignment=CEN
R.cell(prog_row,5,f'=COUNTIF(D{first_item}:D{last_item},"Sim")&" de "&(COUNTA(B{first_item}:B{last_item}))').font=F(9,True)
R.freeze_panes=f"A{hdr_row+1}"
rr+=1
# photo log
R.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
R.cell(rr,1,"LOG FOTOGRÁFICO").font=F(10,True,"FFFFFF"); R.cell(rr,1).fill=fill(C_SEC); R.cell(rr,1).alignment=Alignment(indent=1,vertical="center"); rr+=1
for i,h in enumerate(["Nº","Arquivo · vista / elemento fotografado","Orientação","Hora","Seção FIAMS","Observação"],1):
    c=R.cell(rr,i,h); c.font=F(9,True,C_TITLE); c.fill=fill(C_TH); c.alignment=CEN; c.border=BOX
rr+=1
for i in range(1,31):
    R.cell(rr,1,f"F{i:02d}").font=F(8); R.cell(rr,1).alignment=CEN
    for cc in range(1,7): R.cell(rr,cc).border=BOX
    for cc in range(2,7): R.cell(rr,cc).fill=fill(AMBER_L)
    rr+=1
R.cell(rr+1,1,"Orientação: N, NE, L, SE, S, SO, O, NO (direção para onde a câmera aponta). Nome de arquivo sugerido: [código FIAMS]_F[nº].jpg").font=F(8,False,"5A5A5A",True)
rr+=3
R.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=6)
R.cell(rr,1,"CROQUI DE IMPLANTAÇÃO — desenhar à mão (lote, recuos, acessos, muros, vegetação, norte) e anexar digitalizado ao dossiê (seção 22.2).").font=F(9,True,"8A4B00")
R.cell(rr,1).fill=fill(AMBER); R.cell(rr,1).alignment=Alignment(wrap_text=True,indent=1,vertical="center"); R.row_dimensions[rr].height=28
rr+=1
R.merge_cells(start_row=rr,start_column=1,end_row=rr+14,end_column=6)
R.cell(rr,1).border=BOX
for x in range(rr,rr+15):
    for cc in range(1,7): R.cell(x,cc).border=BOX
R.page_setup.orientation="portrait"; R.page_setup.paperSize=9; R.page_setup.fitToWidth=1; R.page_setup.fitToHeight=0
R.sheet_properties.pageSetUpPr.fitToPage=True; R.print_title_rows=f"{hdr_row}:{hdr_row}"
R.oddFooter.center.text="Roteiro de campo · &P / &N"

# ---------------- MONITORAMENTO ----------------
M=wm; M.sheet_view.showGridLines=False
mw={"A":6,"B":13,"C":20,"D":30,"E":24,"F":24,"G":24,"H":22,"I":13,"J":18,"K":16}
for k,v in mw.items(): M.column_dimensions[k].width=v
M.merge_cells("A1:K1"); M["A1"]="MONITORAMENTO PERIÓDICO — FIAMS (Nível 3)"; M["A1"].font=F(14,True,"FFFFFF"); M["A1"].fill=fill(C_TITLE); M["A1"].alignment=Alignment(vertical="center",indent=1); M.row_dimensions[1].height=30
M.merge_cells("A2:K2"); M["A2"]=f'="Bem: "&IF(FIAMS!{KEY["ID"]}="","—",FIAMS!{KEY["ID"]})&"   ·   "&{DEN}'
M["A2"].font=F(10,True,C_TITLE); M["A2"].fill=fill(C_SUB); M["A2"].alignment=Alignment(indent=1,vertical="center")
M.merge_cells("A3:K3"); M["A3"]="Esta seção transforma a ficha de inventário em instrumento de gestão. Cada linha corresponde a uma inspeção de campo; a última inspeção alimenta automaticamente a aba SIG_Tabela."
M["A3"].font=F(8.5,False,"4A4A4A",True); M["A3"].alignment=Alignment(indent=1,wrap_text=True,vertical="center"); M.row_dimensions[3].height=22
M.merge_cells("A5:K5"); M["A5"]="26. REGISTRO DE INSPEÇÕES"; M["A5"].font=F(11,True,"FFFFFF"); M["A5"].fill=fill(C_SEC); M["A5"].alignment=Alignment(indent=1,vertical="center")
mh=["Nº","Data da inspeção","Situação comparada à última inspeção","Alterações observadas","Novas patologias","Novas intervenções","Alterações no entorno","Necessidade de ação","Próxima inspeção (data)","Responsável","Fotos (nº / pasta)"]
for i,h in enumerate(mh,1):
    c=M.cell(6,i,h); c.font=F(8.5,True,"FFFFFF"); c.fill=fill(C_TITLE); c.alignment=CEN; c.border=BOX
M.row_dimensions[6].height=32
MON_A,MON_B=7,26
for x in range(MON_A,MON_B+1):
    M.cell(x,1,x-MON_A+1).alignment=CEN; M.cell(x,1).font=F(8)
    for cc in range(1,12):
        c=M.cell(x,cc); c.border=BOX; c.alignment=WRAP; c.font=F(9) if cc>1 else F(8)
        if cc>1: c.fill=fill(AMBER if cc in (2,3,4,5,6,7) else "FFFFFF")
    M.cell(x,2).number_format="DD/MM/YYYY"; M.cell(x,9).number_format="DD/MM/YYYY"
    add_dv(M,"SITMON",f"C{x}"); add_dv(M,"ACAO",f"H{x}")
    M.row_dimensions[x].height=30
M.cell(MON_B+1,2,"Prazo de reinspeção sugerido: 12 meses (Classes A/B: 6 meses; risco Alto/Crítico: 3 meses).").font=F(8,False,"5A5A5A",True)
mr=MON_B+3
M.merge_cells(start_row=mr,start_column=1,end_row=mr,end_column=11)
M.cell(mr,1,"27. MATRIZ DE MONITORAMENTO (escala 1 a 5 — 5 = situação mais favorável ao bem)").font=F(11,True,"FFFFFF"); M.cell(mr,1).fill=fill(C_SEC); M.cell(mr,1).alignment=Alignment(indent=1,vertical="center")
mr+=1
mh2=["Nº","Indicador","","Leitura da escala (1 → 5)","","Situação inicial","Situação atual","Variação","Tendência","Observação",""]
for i,h in enumerate(mh2,1):
    c=M.cell(mr,i,h); c.font=F(8.5,True,"FFFFFF"); c.fill=fill(C_TITLE); c.alignment=CEN; c.border=BOX
M.merge_cells(start_row=mr,start_column=2,end_row=mr,end_column=3); M.merge_cells(start_row=mr,start_column=4,end_row=mr,end_column=5); M.merge_cells(start_row=mr,start_column=10,end_row=mr,end_column=11)
mr+=1
inds=[("Integridade","1 muito baixa → 5 muito alta"),("Autenticidade","1 muito baixa → 5 muito alta"),("Estado de conservação","1 ruína → 5 excelente"),("Integridade do entorno","1 descaracterizado → 5 preservado"),("Compatibilidade de uso","1 incompatível → 5 compatível"),("Risco de descaracterização","1 crítico → 5 muito baixo"),("Risco de perda","1 crítico → 5 muito baixo"),("Pressão imobiliária","1 muito alta → 5 muito baixa"),("Proteção legal","1 nenhuma → 5 tombamento"),("Documentação disponível","1 inexistente → 5 completa")]
ind_a=mr
dvn=DataValidation(type="whole",operator="between",formula1="1",formula2="5",allow_blank=True,showErrorMessage=True,error="Informe um inteiro de 1 a 5."); M.add_data_validation(dvn)
for i,(ind,esc) in enumerate(inds,1):
    M.cell(mr,1,i).alignment=CEN
    M.merge_cells(start_row=mr,start_column=2,end_row=mr,end_column=3); M.cell(mr,2,ind).font=F(9,True)
    M.merge_cells(start_row=mr,start_column=4,end_row=mr,end_column=5); M.cell(mr,4,esc).font=F(8,False,"5A5A5A",True)
    for cc in (6,7): M.cell(mr,cc).fill=fill(AMBER); M.cell(mr,cc).alignment=CEN; dvn.add(f"{get_column_letter(cc)}{mr}")
    M.cell(mr,8,f'=IF(OR(F{mr}="",G{mr}=""),"",G{mr}-F{mr})').alignment=CEN
    M.cell(mr,9,f'=IF(H{mr}="","",IF(H{mr}>0,"↑ melhora",IF(H{mr}<0,"↓ piora","→ estável")))').alignment=CEN
    M.merge_cells(start_row=mr,start_column=10,end_row=mr,end_column=11)
    for cc in range(1,12): M.cell(mr,cc).border=BOX
    M.row_dimensions[mr].height=18; mr+=1
ind_b=mr-1
M.merge_cells(start_row=mr,start_column=2,end_row=mr,end_column=5); M.cell(mr,2,"Índice médio de monitoramento").font=F(9,True,C_TITLE)
for cc,L in ((6,"F"),(7,"G")):
    M.cell(mr,cc,f'=IF(COUNT({L}{ind_a}:{L}{ind_b})=0,"",AVERAGE({L}{ind_a}:{L}{ind_b}))').number_format="0.0"
    M.cell(mr,cc).font=F(10,True); M.cell(mr,cc).alignment=CEN
M.cell(mr,8,f'=IF(OR(F{mr}="",G{mr}=""),"",G{mr}-F{mr})').number_format="+0.0;-0.0;0.0"; M.cell(mr,8).alignment=CEN
M.cell(mr,9,f'=IF(H{mr}="","",IF(H{mr}>0,"↑ melhora",IF(H{mr}<0,"↓ piora","→ estável")))').alignment=CEN
for cc in range(1,12): M.cell(mr,cc).border=BOX; M.cell(mr,cc).fill=fill(C_SUB)
M.cell(mr+1,2,"Situação inicial = primeiro registro FIAMS; situação atual = última inspeção. As variações são calculadas automaticamente.").font=F(8,False,"5A5A5A",True)
M.freeze_panes="A7"
M.page_setup.orientation="landscape"; M.page_setup.paperSize=9; M.page_setup.fitToWidth=1; M.page_setup.fitToHeight=0; M.sheet_properties.pageSetUpPr.fitToPage=True

Bdate=f"Monitoramento!$B${MON_A}:$B${MON_B}"
def latest(colL): return f'IF(COUNT({Bdate})=0,"",IFERROR(INDEX(Monitoramento!${colL}${MON_A}:${colL}${MON_B},MATCH(MAX({Bdate}),{Bdate},0))&"",""))'

# ---------------- SIG ----------------
def link(k,num=False):
    ref=f"FIAMS!{KEY[k]}"
    return f'=IF({ref}="","",{ref})'
sig=[("ID_FIAMS","Texto","Código único do bem","1.1.1",link("ID")),
("NOME_ATUAL","Texto","Nome atual","1.2.1",link("NOME_ATUAL")),
("NOME_ORIGINAL","Texto","Nome original","1.2.2",link("NOME_ORIG")),
("ENDERECO","Texto","Logradouro, número","1.3.1–1.3.2",f'=TRIM(FIAMS!{KEY["LOGR"]}&IF(FIAMS!{KEY["NUM"]}="","",", "&FIAMS!{KEY["NUM"]}))'),
("BAIRRO","Texto","Bairro","1.3.4",link("BAIRRO")),
("LATITUDE","Real (graus decimais, SIRGAS 2000)","Latitude","1.4.1",link("LAT")),
("LONGITUDE","Real (graus decimais, SIRGAS 2000)","Longitude","1.4.2",link("LON")),
("ANO_PROJETO","Texto/Inteiro","Data do projeto","4 — Projeto",link("ANO_PROJ")),
("ANO_CONCLUSAO","Texto/Inteiro","Data de conclusão","4 — Conclusão",link("ANO_CONC")),
("ARQUITETO","Texto","Autor do projeto","3.1.1",link("ARQ")),
("TIPOLOGIA","Texto","Tipologia original","2.2.1",link("TIPO")),
("LINGUAGEM_MODERNA","Texto (domínio)","Classificação principal","2.3.1",link("LING")),
("USO_ORIGINAL","Texto","Uso original","11.1",link("USO_ORIG")),
("USO_ATUAL","Texto","Uso atual","11.2",link("USO_ATUAL")),
("PAVIMENTOS","Inteiro","Número de pavimentos","8.1.1",link("PAV")),
("AREA_LOTE","Real (m²)","Área do lote","7.1.1",link("AREA_LOTE")),
("AREA_CONSTRUIDA","Real (m²)","Área construída aproximada","8.1.3",link("AREA_CONST")),
("PROTECAO","Texto (domínio)","Situação de proteção","1.5.1",link("PROT")),
("INTEGRIDADE","Texto (domínio)","Grau geral de integridade","14.1.10",link("INTEG")),
("IIM_PCT","Real (0–1)","Índice de Integridade Modernista","14.2.2",link("IIM")),
("IIM_CATEGORIA","Texto (domínio)","Categoria do IIM","14.2.3",link("IIM_CAT")),
("AUTENTICIDADE","Texto (domínio)","Autenticidade (síntese)","20.4",link("AUTENT")),
("CONSERVACAO","Texto (domínio)","Estado de conservação geral","16.1",link("CONS")),
("VALOR_HISTORICO","Texto (domínio)","Valor histórico","19.1",link("V_HIST")),
("VALOR_ARQUITETONICO","Texto (domínio)","Valor arquitetônico","19.2",link("V_ARQ")),
("VALOR_URBANISTICO","Texto (domínio)","Valor urbanístico","19.3",link("V_URB")),
("VALOR_SOCIAL","Texto (domínio)","Valor social","19.6",link("V_SOC")),
("VALOR_DOCUMENTAL","Texto (domínio)","Valor documental","19.8",link("V_DOC")),
("SIGNIFICANCIA","Texto (domínio)","Representatividade do modernismo maranhense","19.9",link("SIGN")),
("RISCO","Texto (domínio)","Classificação de risco","18.1",link("RISCO")),
("PRIORIDADE","Texto (domínio)","Classe de prioridade","21.1",link("PRIOR")),
("ULTIMA_VISTORIA","Data","Data da inspeção mais recente","Monitoramento — 26",f'=IF(COUNT({Bdate})=0,"",MAX({Bdate}))'),
("PROXIMA_VISTORIA","Data","Próxima inspeção prevista (da inspeção mais recente)","Monitoramento — 26",f'=IF(COUNT({Bdate})=0,"",IFERROR(INDEX(Monitoramento!$I${MON_A}:$I${MON_B},MATCH(MAX({Bdate}),{Bdate},0)),""))'),
("STATUS","Texto (domínio)","Situação na inspeção mais recente","Monitoramento — 26","="+latest("C")),
("LINK_DOSSIE","Texto (URL/caminho)","Pasta do dossiê gráfico","22.0",link("LINK")),
("FOTOGRAFIA_PRINCIPAL","Texto (arquivo)","Fotografia principal","22.3.1",link("FOTO")),
]
S=wsig; S.sheet_view.showGridLines=False
for i,(a,t,d,o,fx) in enumerate(sig,1):
    c=S.cell(1,i,a); c.font=F(9,True,"FFFFFF"); c.fill=fill(C_TITLE); c.alignment=CEN; c.border=BOX
    v=S.cell(2,i,fx); v.font=F(9); v.border=BOX; v.alignment=WRAP
    if a in ("LATITUDE","LONGITUDE"): v.number_format="0.000000"
    if a=="IIM_PCT": v.number_format="0.0%"
    if a in ("ULTIMA_VISTORIA","PROXIMA_VISTORIA"): v.number_format="DD/MM/YYYY"
    S.column_dimensions[get_column_letter(i)].width=max(14,len(a)+3)
S.row_dimensions[1].height=22; S.freeze_panes="B2"
D_=wdic; D_.sheet_view.showGridLines=False
D_.merge_cells("A1:E1"); D_["A1"]="DICIONÁRIO DE DADOS — TABELA DE ATRIBUTOS PARA SIG (QGIS / plataforma web)"; D_["A1"].font=F(12,True,"FFFFFF"); D_["A1"].fill=fill(C_TITLE); D_["A1"].alignment=Alignment(indent=1,vertical="center"); D_.row_dimensions[1].height=26
D_.merge_cells("A2:E2"); D_["A2"]="A aba SIG_Tabela é preenchida automaticamente a partir da ficha. Para compor a camada do inventário, exporte a linha 2 de cada ficha (CSV UTF-8) e una em uma única tabela; use LATITUDE/LONGITUDE (EPSG:4674 — SIRGAS 2000) para criar a camada de pontos."
D_["A2"].font=F(8.5,False,"4A4A4A",True); D_["A2"].alignment=Alignment(wrap_text=True,indent=1,vertical="center"); D_.row_dimensions[2].height=32
for i,h in enumerate(["Atributo","Tipo","Descrição","Origem na FIAMS","Coleta"],1):
    c=D_.cell(4,i,h); c.font=F(9,True,"FFFFFF"); c.fill=fill(C_SEC); c.alignment=CEN; c.border=BOX
colmap={"ID_FIAMS":DOC,"NOME_ATUAL":CAMPO,"NOME_ORIGINAL":DOC,"ENDERECO":CD,"BAIRRO":CD,"LATITUDE":CAMPO,"LONGITUDE":CAMPO,"PAVIMENTOS":CAMPO,"AREA_LOTE":CD,"AREA_CONSTRUIDA":CD,"USO_ATUAL":CAMPO,"CONSERVACAO":CAMPO,"FOTOGRAFIA_PRINCIPAL":CAMPO,"ULTIMA_VISTORIA":CAMPO,"STATUS":CAMPO,"PROXIMA_VISTORIA":ANA,"IIM_PCT":ANA,"IIM_CATEGORIA":ANA}
for i,(a,t,d,o,fx) in enumerate(sig,5):
    for j,v in enumerate([a,t,d,o],1):
        c=D_.cell(i,j,v); c.font=F(9,j==1); c.border=BOX; c.alignment=WRAP
    tg=colmap.get(a, DOC if o[:1] in "1234" and a not in ("LINGUAGEM_MODERNA",) and not o.startswith(("14","16","18","19","20","21")) else ANA)
    c=D_.cell(i,5,tg); c.font=F(7.5,True,"FFFFFF"); c.fill=fill(TAGFILL[tg]); c.alignment=CEN; c.border=BOX
for k,v in {"A":24,"B":26,"C":42,"D":20,"E":13}.items(): D_.column_dimensions[k].width=v

# ---------------- APRESENTAÇÃO ----------------
A=ws0; A.sheet_view.showGridLines=False
A.column_dimensions["A"].width=3; A.column_dimensions["B"].width=30; A.column_dimensions["C"].width=95
A.merge_cells("B2:C2"); A["B2"]="FIAMS"; A["B2"].font=F(26,True,"FFFFFF"); A["B2"].fill=fill(C_TITLE); A["B2"].alignment=Alignment(indent=1,vertical="center"); A.row_dimensions[2].height=44
A.merge_cells("B3:C3"); A["B3"]="Ficha de Inventário da Arquitetura Moderna de São Luís (1930–1980)"; A["B3"].font=F(13,True,"FFFFFF"); A["B3"].fill=fill(C_TITLE); A["B3"].alignment=Alignment(indent=1,vertical="center"); A.row_dimensions[3].height=24
A.merge_cells("B4:C4"); A["B4"]="Inventário de Identificação, Documentação, Avaliação e Monitoramento do Acervo Modernista de São Luís"; A["B4"].font=F(10,False,"FFFFFF",True); A["B4"].fill=fill(C_SEC); A["B4"].alignment=Alignment(indent=1,vertical="center"); A.row_dimensions[4].height=20
ar=6
def ahead(t):
    global ar
    A.merge_cells(start_row=ar,start_column=2,end_row=ar,end_column=3)
    c=A.cell(ar,2,t); c.font=F(10.5,True,C_TITLE); c.fill=fill(C_SUB); c.alignment=Alignment(indent=1,vertical="center"); A.row_dimensions[ar].height=20; ar+=1
def arow(k,v,h=None,kfill=None,kfont=None):
    global ar
    a=A.cell(ar,2,k); a.font=kfont or F(9,True); a.alignment=Alignment(wrap_text=True,vertical="top",indent=1)
    if kfill: a.fill=fill(kfill)
    b=A.cell(ar,3,v); b.font=F(9); b.alignment=Alignment(wrap_text=True,vertical="top")
    A.row_dimensions[ar].height=h or max(16,14*(len(v)//105+1)); ar+=1
ahead("IDENTIFICAÇÃO DO INSTRUMENTO")
arow("Instrumento","Ficha individual de bem imóvel")
arow("Natureza","Inventário temático de arquitetura moderna")
arow("Recorte temporal","1930–1980")
arow("Recorte territorial","Bairros de Monte Castelo (Retiro Natal, Fátima, Vila Passos, Belira, Canto da Fabril, Apeadouro, Alemanha, Coreia, Vila Ivar Saldanha), João Paulo e Filipinho, no município de São Luís, Maranhão")
arow("Base metodológica","INBI/INBI-SU/IPHAN + SICG/IPHAN + DOCOMOMO + experiências brasileiras de inventariação da arquitetura moderna analisadas por Guedes")
arow("Finalidade","Identificação, documentação, interpretação, avaliação patrimonial e monitoramento das transformações do acervo moderno")
arow("Instituição","LUPA — Laboratório de Urbanismo, Paisagem, Arquitetura e Artes · Centro Universitário UNDB")
ar+=1
ahead("LEGENDA — ORIGEM DA INFORMAÇÃO")
for tg,desc in ((CAMPO,"Registro obrigatório na visita in loco: observar, medir, fotografar ou perguntar. Células de resposta em laranja-claro na aba FIAMS e reunidas na aba Roteiro_Campo."),(CD,"Informação documental que DEVE ser confirmada ou corrigida em campo (endereço, dimensões do lote, datas de alteração). Também destacada em laranja."),(DOC,"Pesquisa em arquivo, cartório, cadastro municipal, bibliografia, cartografia e iconografia — preferencialmente antes da visita."),(ANA,"Avaliação técnica realizada após o campo, com base no conjunto de dados (autenticidade, integridade, valores, prioridade).")):
    arow(tg,desc,kfill=TAGFILL[tg],kfont=F(9,True,"FFFFFF"))
ar+=1
ahead("ESTRUTURA DA PASTA DE TRABALHO")
for k,v in (("Apresentação","Identificação do instrumento, legenda, instruções e fundamentação metodológica."),("FIAMS","Nível 1 — Ficha de inventário, seções 1 a 28. Painel de preenchimento no topo (percentual por tipo de coleta e IIM)."),("Roteiro_Campo","Checklist operacional da visita, gerado a partir de todos os itens de CAMPO da ficha, com links diretos para as células de resposta, equipamentos, log fotográfico e área de croqui. Imprimir em A4 retrato."),("Monitoramento","Nível 3 — Registro periódico de inspeções (seção 26) e matriz quantitativa de monitoramento com cálculo automático de tendência (seção 27)."),("SIG_Tabela","Seção 29 — Linha de atributos vinculada por fórmulas à ficha, pronta para exportação a QGIS / plataforma web."),("SIG_Dicionário","Definição de cada atributo, tipo de dado e seção de origem.")):
    arow(k,v)
ar+=1
ahead("INSTRUÇÕES DE PREENCHIMENTO")
ins=["Salve uma cópia desta pasta por bem, nomeando o arquivo com o código FIAMS (ex.: FIAMS-MC-001.xlsx).",
"Antes da visita, preencha os campos DOCUMENTAL disponíveis (seções 1, 3, 4, 5 e 6) e imprima a aba Roteiro_Campo.",
"Na visita, percorra o roteiro na ordem: exterior e entorno → lote e implantação → volumetria e fachadas → interior (se houver acesso) → entrevista. Registre as respostas nas células laranja da ficha.",
"Use as listas suspensas sempre que existirem; elas padronizam os dados para o SIG. Informação complementar vai na coluna 'Fonte · grau de certeza'.",
"Não resolva divergências entre fontes silenciosamente: registre ambas as versões na coluna E e indique a fonte de cada uma.",
"Quando o atributo não puder ser verificado (acesso negado, elemento encoberto), marque 'Não verificável' / 's.d.' em vez de deixar em branco.",
"Na matriz A1–A7 (14.2), use 'n.a.' apenas para atributos comprovadamente inexistentes no projeto original. O IIM é calculado automaticamente.",
"Após o campo, complete as seções de ANÁLISE (13, 14, 19, 20, 21, 28) e atualize a versão e a data da ficha (1.1).",
"Cada nova visita de acompanhamento deve ser lançada como nova linha na aba Monitoramento, sem sobrescrever registros anteriores."]
for i,t in enumerate(ins,1): arow(f"{i}.",t)
ar+=1
ahead("ARQUITETURA METODOLÓGICA EM TRÊS NÍVEIS")
arow("Nível 1 — Ficha de inventário","Aplicada aos bens confirmados: identificação, localização, datação, autoria, tipologia, características modernas, uso, coordenadas, história, implantação, arquitetura, técnica construtiva, entorno, alterações, conservação, autenticidade, integridade e valores.")
arow("Nível 2 — Dossiê documental","Fotografias históricas e atuais, plantas, cortes, fachadas, redesenhos, mapas, documentos, entrevistas e referências.")
arow("Nível 3 — Ficha de monitoramento","Atualizada periodicamente: estado de conservação, alterações, ameaças, intervenções, mudanças no entorno, atualização fotográfica e recomendações.")
ar+=1
ahead("FUNDAMENTAÇÃO")
arow("Justificativa","A análise de quinze inventários brasileiros realizada por Guedes evidencia grande diversidade de instrumentos e recorrência de lacunas: fotografias históricas, desenhos técnicos, análise formal e técnico-construtiva, relação edifício–lote–entorno, alterações, valor/significância, instruções de preenchimento, identificação do relator e estado de proteção/conservação. A estrutura da FIAMS responde a cada uma delas.",h=58)
arow("INBI-SU / IPHAN","Base territorial e documental: pesquisa histórica, levantamento físico-arquitetônico e entrevistas, articulando lote, edificação, uso e contexto urbano. O próprio IPHAN realizou o INBI-SU de São Luís, com levantamentos planialtimétricos, físico-arquitetônicos e entrevistas organizados em banco de dados.",h=44)
arow("SICG / IPHAN","Identificação, localização, dados do bem, estado de conservação/preservação, entorno, síntese histórica, multimídia, documentos e proteções — campos que Guedes demonstra ser necessário complementar para a especificidade moderna.",h=32)
arow("DOCOMOMO","Especificidade da arquitetura moderna: identificação, história, descrição, condições físicas, potencial de conservação, avaliações técnica, social e cultural-estética, valores históricos/de referência e registro das alterações.",h=32)
arow("ICOMOS / UNESCO","Integridade, autenticidade, fatores de ameaça, proteção, conservação, gerenciamento e monitoramento.")
arow("Matriz A1–A7 e IIM","Atributos definidores da integridade modernista e Índice de Integridade Modernista conforme a metodologia do projeto (Trindade e Silva; Lopes, 2026).")
arow("Critério de explicitação","Não basta reproduzir modelos de fichas: é necessário explicitar objetivos, critérios de seleção, metodologia e a relação entre fundamentação teórica e campos efetivamente preenchidos — deficiência recorrente nos inventários brasileiros apontada por Guedes.",h=32)
arow("Denominação","A FIAMS é uma adaptação temática e crítica dos referenciais INBI/INBI-SU e SICG/IPHAN, complementados pelas fichas DOCOMOMO e pelas experiências brasileiras de inventariação da arquitetura moderna — e não uma reprodução de formulário oficial do IPHAN.",h=32)
arow("Referência a completar","GUEDES, [nome]. [Título da tese]. [Programa, instituição, ano]. — completar conforme ABNT NBR 6023.",kfill=AMBER)
A.page_setup.orientation="portrait"; A.page_setup.paperSize=9; A.page_setup.fitToWidth=1; A.page_setup.fitToHeight=0; A.sheet_properties.pageSetUpPr.fitToPage=True

# default font for untouched cells
for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if c.value is not None and c.font.name!="Arial":
                c.font=F(9)
wb.active=1
import sys
from pathlib import Path
HERE=Path(__file__).resolve().parent; WORK=HERE/"work"; WORK.mkdir(exist_ok=True)
out=sys.argv[1] if len(sys.argv)>1 else str(WORK/"template.xlsx")
wb.save(out); print("ficha:",out,LAST,len(FIELD_ITEMS))
import json; json.dump({"schema":SC,"lists":LISTS},open(WORK/"schema_raw.json","w"),ensure_ascii=False,indent=0); print("schema:",len(SC),"elementos")
