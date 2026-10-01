import json, glob, datetime, re
from pathlib import Path
HERE=Path(__file__).resolve().parent; WORK=HERE/'work'; INPUT=HERE/'input'; ROOT=HERE.parent
from openpyxl import load_workbook
rowmap=json.load(open(WORK/'rowmap.json')); st={s['code']:s for s in json.load(open(WORK/'stats.json'))}
sch=json.loads(open(ROOT/'site'/'campo'/'js'/'schema.js').read().split("export const SCHEMA=")[1].split(";\nexport const LISTS=")[0])
T=load_workbook(str(WORK/'template.xlsx'))['FIAMS']
inv=load_workbook(str(INPUT/'Inventario_MonteCastelo_JoaoPaulo_Filipinho_v4_GERAL_1.xlsx'))['Consolidado']
coords={r[0]:(r[11],r[12],r[6]) for r in inv.iter_rows(min_row=2,values_only=True) if r[0]}
def cv(v):
    if isinstance(v,(datetime.datetime,datetime.date)): return v.date().isoformat() if isinstance(v,datetime.datetime) else v.isoformat()
    if isinstance(v,str): return v.strip()
    return v
tables={}
for s in sch['sections']:
    for it in s['items']:
        if it['type'] in ('table','rtable'): tables[it['id']]=it
out=[]
for f in sorted(glob.glob(str(WORK/'fichas'/'FIAMS-*.xlsx'))):
    ws=load_workbook(f)['FIAMS']; d={}; pre=[]
    def setv(key,r,c):
        v=ws.cell(r,c).value
        if v in (None,"") or (isinstance(v,str) and v.startswith("=")): return
        d[key]=cv(v)
        if v!=T.cell(r,c).value: pre.append(key)
    for key,row in rowmap.items():
        if key in tables:
            r0,n=row; t=tables[key]
            if t['type']=='table':
                for i in range(n):
                    for c,col in ((4,'D'),(5,'E'),(6,'F')):
                        if key=='iim' and col=='E': continue
                        setv(f"{key}|{i}|{col}",r0+i,c)
            else:
                rows=[]
                for i in range(n):
                    rr={}
                    for c,col in ((2,'B'),(4,'D'),(5,'E'),(6,'F')):
                        v=ws.cell(r0+i,c).value
                        if v not in (None,""): rr[col]=cv(v)
                    if rr: rows.append(rr)
                if rows: d[key]=rows; pre.append(key)
        else:
            setv(key,row,4); setv(key+"|fonte",row,5)
    code=ws.cell(rowmap["1.1.1"],4).value; s=st[code]; lat,lon,loc=coords[s['n']]
    d["__pre"]=sorted(set(pre))
    out.append({"id":code,"n":s['n'],"nome":s['nome'],"endereco":s['end'],"bairro":s['aba'],"localidade":loc,"lat":lat,"lon":lon,"autor":s['autor'],
        "data_ref":str(s['data']) if s['data'] not in (None,'-') else None,"funcao":s['funcao'],"origem":s['origem'],"levantamento_2026":s['surv'] or None,"alertas":s['alerts'] or None,"seed":d})
json.dump(out,open(ROOT/'site'/'campo'/'data'/'seed.json','w'),ensure_ascii=False)
# SQL
def q(v):
    if v is None: return "null"
    if isinstance(v,(int,float)): return repr(v)
    return "'"+str(v).replace("'","''")+"'"
lines=["-- FIAMS · carga inicial dos 86 imóveis do Inventário v4 (fichas pré-preenchidas)","-- Execute DEPOIS de 01_schema.sql. Pode ser executado de novo: atualiza os registros existentes.",
"insert into public.imoveis (id,n,nome,endereco,bairro,localidade,lat,lon,autor,data_ref,funcao,origem,levantamento_2026,alertas,seed) values"]
vals=[]
for o in out:
    js=json.dumps(o['seed'],ensure_ascii=False)
    assert "$fiams$" not in js
    vals.append(f"({q(o['id'])},{o['n']},{q(o['nome'])},{q(o['endereco'])},{q(o['bairro'])},{q(o['localidade'])},{q(o['lat'])},{q(o['lon'])},{q(o['autor'])},{q(o['data_ref'])},{q(o['funcao'])},{q(o['origem'])},{q(o['levantamento_2026'])},{q(o['alertas'])},$fiams${js}$fiams$::jsonb)")
lines.append(",\n".join(vals))
lines.append("on conflict (id) do update set n=excluded.n,nome=excluded.nome,endereco=excluded.endereco,bairro=excluded.bairro,localidade=excluded.localidade,lat=excluded.lat,lon=excluded.lon,autor=excluded.autor,data_ref=excluded.data_ref,funcao=excluded.funcao,origem=excluded.origem,levantamento_2026=excluded.levantamento_2026,alertas=excluded.alertas,seed=excluded.seed;")
open(ROOT/'supabase'/'02_seed_imoveis.sql','w').write("\n".join(lines)+"\n")
print(len(out),'imóveis;',sum(len(o['seed']) for o in out),'valores pré-preenchidos')
