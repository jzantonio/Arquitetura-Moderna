import json
from pathlib import Path
WORK=Path(__file__).resolve().parent/'work'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
st=json.load(open(WORK/'stats.json'))
F=lambda sz=9,b=False,c="1A1A1A",u=None:Font(name="Arial",size=sz,bold=b,color=c,underline=u)
fill=lambda c:PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color="A6B4BE"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); ws=wb.active; ws.title="Índice"; ws.sheet_view.showGridLines=False
cols=[("Código FIAMS",16),("Denominação",42),("Endereço",32),("Bairro (inventário)",16),("Aba",14),("Autor",28),("Data",18),("Função",18),("Origem do registro",18),("Levantamento 2026 (Trindade e Silva; Lopes)",30),("Campos pré-preenchidos",12),("Alertas para o trabalho de campo",40),("Arquivo",14)]
ws.merge_cells("A1:M1"); ws["A1"]="FIAMS — ÍNDICE DAS FICHAS PRÉ-PREENCHIDAS (Inventário v4)"; ws["A1"].font=F(14,True,"FFFFFF"); ws["A1"].fill=fill("1F3A4D"); ws["A1"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[1].height=30
ws.merge_cells("A2:M2"); ws["A2"]="86 fichas geradas a partir da aba Consolidado. Os dados pré-preenchidos aparecem em azul em cada ficha; campos em branco ficam para a equipe. Clique em 'abrir' para abrir a ficha (mantenha este índice na mesma pasta das fichas)."
ws["A2"].font=F(9,False,"FFFFFF"); ws["A2"].fill=fill("2E5A6E"); ws["A2"].alignment=Alignment(wrap_text=True,vertical="center",indent=1); ws.row_dimensions[2].height=28
H=5
for i,(h,w) in enumerate(cols,1):
    c=ws.cell(H,i,h); c.font=F(9,True,"FFFFFF"); c.fill=fill("1F3A4D"); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=BOX
    ws.column_dimensions[get_column_letter(i)].width=w
ws.row_dimensions[H].height=32
for k,s in enumerate(st):
    r=H+1+k
    vals=[s['code'],s['nome'],s['end'],s['bairro'],s['aba'],s['autor'],s['data'] if s['data'] not in (None,'-') else '',s['funcao'],s['origem'],s['surv'],s['filled'],s['alerts'],"abrir"]
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.font=F(8.5); c.border=BOX; c.alignment=Alignment(wrap_text=True,vertical="top")
    ws.cell(r,11).alignment=Alignment(horizontal="center",vertical="top")
    ln=ws.cell(r,13); ln.hyperlink=s['file']; ln.font=F(8.5,False,"1F5FA8","single")
    if s['alerts']: ws.cell(r,12).fill=fill("FFE9C7")
last=H+len(st)
ws.auto_filter.ref=f"A{H}:M{last}"; ws.freeze_panes=f"C{H+1}"
# resumo
ws["A3"]="Resumo"; ws["A3"].font=F(9,True)
items=[("Fichas",f"=COUNTA(A{H+1}:A{last})"),("Com alertas",f'=COUNTIF(L{H+1}:L{last},"?*")'),("Endereço a confirmar",f'=COUNTIF(L{H+1}:L{last},"*endereço*")'),("Fora do recorte temporal",f'=COUNTIF(L{H+1}:L{last},"*fora do recorte*")'),("Demolidos/substituídos",f'=COUNTIF(L{H+1}:L{last},"*demolido*")'),("Função inferida",f'=COUNTIF(L{H+1}:L{last},"*função inferida*")')]
col=2
for lab,fx in items:
    ws.cell(3,col,lab).font=F(8,True,"1F3A4D"); c=ws.cell(4,col,fx); c.font=F(10,True); c.alignment=Alignment(horizontal="left")
    col+=2 if col<11 else 1
ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=9; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_title_rows=f"{H}:{H}"
wb.save(str(WORK/"fichas"/"00_Indice_Fichas_FIAMS.xlsx"))
