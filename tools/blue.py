import glob, copy
from pathlib import Path
WORK=Path(__file__).resolve().parent/'work'
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
T=load_workbook(str(WORK/'template.xlsx'))['FIAMS']
tv={(r,c):T.cell(r,c).value for r in range(1,T.max_row+1) for c in range(2,7)}
for f in sorted(glob.glob(str(WORK/'fichas'/'FIAMS-*.xlsx'))):
    wb=load_workbook(f); ws=wb['FIAMS']
    for r in range(6,ws.max_row+1):
        for c in range(2,7):
            cell=ws.cell(r,c)
            if cell.value not in (None,"") and cell.value!=tv.get((r,c)) and not str(cell.value).startswith("="):
                ft=copy.copy(cell.font); cell.font=Font(name="Arial",size=ft.size,bold=ft.bold,italic=ft.italic,color="1F5FA8")
    A=wb['Apresentação']; ar=A.max_row+2
    A.merge_cells(start_row=ar,start_column=2,end_row=ar,end_column=3)
    c=A.cell(ar,2,"PRÉ-PREENCHIMENTO A PARTIR DO INVENTÁRIO v4"); c.font=Font(name="Arial",size=10.5,bold=True,color="1F3A4D"); c.fill=PatternFill("solid",start_color="DCE6EC",end_color="DCE6EC")
    txt=("Esta ficha (versão v0.1) foi pré-preenchida com os dados já existentes na planilha Inventario_MonteCastelo_JoaoPaulo_Filipinho_v4_GERAL. "
         "Os valores pré-preenchidos aparecem em AZUL; a coluna 'Fonte · grau de certeza' indica de qual campo da planilha cada dado foi extraído. "
         "Campos em branco não têm informação no inventário e devem ser preenchidos pela equipe. Dados de campo pré-preenchidos (laranja) vêm da bibliografia ou do levantamento de 2026 e DEVEM ser verificados in loco; "
         "classificações analíticas marcadas 'validar' são preliminares. Ao revisar, troque a cor da fonte para preto e atualize a versão (1.1.6).")
    A.cell(ar+1,2,"Como ler").font=Font(name="Arial",size=9,bold=True)
    t=A.cell(ar+1,3,txt); t.font=Font(name="Arial",size=9); t.alignment=Alignment(wrap_text=True,vertical="top"); A.row_dimensions[ar+1].height=80
    wb.save(f)
print("ok")
